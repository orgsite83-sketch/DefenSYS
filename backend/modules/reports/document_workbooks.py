"""Editable, printable report worksheets with context once above their matrix."""
import re
from datetime import date

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def unique_sheet_name(workbook, label):
    label = re.sub(r'[\\/*?:\[\]]', '-', str(label)).strip(" '") or 'Report'
    candidate, suffix = label[:31], 1
    while candidate.lower() in {name.lower() for name in workbook.sheetnames}:
        suffix += 1
        candidate = label[:31 - len(str(suffix)) - 1] + f' {suffix}'
    return candidate


def text_parts(text, limit=1000):
    """Keep every character while avoiding word breaks between worksheet rows."""
    while len(text) > limit:
        spaces = list(re.finditer(r'\s+', text[:limit]))
        end = spaces[-1].end() if spaces and spaces[-1].end() > limit // 2 else limit
        yield text[:end]
        text = text[end:]
    if text:
        yield text


def write_document_sheet(workbook, spec, generated_by, include_signatures, signatories):
    from reports.export_formatters import _write_xlsx_signatories, _xlsx_value
    from reports.workbook_preview import column_points, display_value, font_name
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph
    from xml.sax.saxutils import escape
    sheet = workbook.create_sheet(unique_sheet_name(workbook, spec['name']))
    sheet.sheet_view.showGridLines = True
    sheet.sheet_view.zoomScale = 100
    columns = spec['columns']
    count = len(columns)
    edge = Side(style='thin', color='CBD5E1')
    border = Border(left=edge, right=edge, top=edge, bottom=edge)
    font = Font(name='Times New Roman', size=10, color='000000')
    heading = Font(name='Times New Roman', size=10, bold=True, color='000000')
    total_width = sum(col.get('width', 16) for col in columns)
    for index, col in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = col.get('width', 16)

    def fitted_height(cell, first, last):
        size = cell.font.sz or 10
        paragraph = Paragraph(escape(display_value(cell)).replace('\n', '<br/>') or '&#160;',
                              ParagraphStyle('Fit', fontName=font_name(cell.font), fontSize=size, leading=size * 1.15))
        width = sum(column_points(sheet, col) for col in range(first, last + 1)) - 6
        return paragraph.wrap(max(8, width), 10000)[1] + 8

    def merge_text(row, value, *, start=1, end=None, bold=False, size=10, centered=False, fill=False):
        end = end or count
        for col in range(start, end + 1):
            sheet.cell(row, col).border = border if fill else Border()
        if end > start:
            sheet.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
        cell = sheet.cell(row, start, str(value or ''))
        cell.data_type = 's'
        cell.font = Font(name='Times New Roman', size=size, bold=bold, color='000000')
        cell.alignment = Alignment(horizontal='center' if centered else 'left', vertical='center' if centered or fill else 'top', wrap_text=True)
        if fill:
            cell.fill = PatternFill('solid', fgColor='F8FAFC')
        sheet.row_dimensions[row].height = min(409, max(20, fitted_height(cell, start, end)))

    merge_text(1, spec['title'], bold=True, size=14, centered=True)
    merge_text(2, spec.get('subtitle', ''), bold=True, centered=True)
    sheet.row_dimensions[3].height = 8
    row = 4
    label_end = max(1, min(count - 1, round(count * .28)))
    for item in spec.get('metadata', []):
        merge_text(row, item['label'], end=label_end, bold=True, fill=True)
        label_height = sheet.row_dimensions[row].height
        merge_text(row, item.get('value', ''), start=label_end + 1, fill=True)
        sheet.row_dimensions[row].height = max(label_height, sheet.row_dimensions[row].height)
        row += 1
    sheet.row_dimensions[row].height = 8
    row += 1
    merge_text(row, spec.get('section_title', 'Report Data'), bold=True, size=12)
    row += 1
    if spec.get('caption'):
        merge_text(row, spec['caption'], size=10)
        row += 1
    header_row = row
    for index, col in enumerate(columns, 1):
        cell = sheet.cell(row, index, col['label'])
        cell.data_type = 's'
        cell.font = heading
        cell.fill = PatternFill('solid', fgColor='F1F5F9')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border
    sheet.row_dimensions[row].height = max(32, max(fitted_height(sheet.cell(row, index), index, index) for index in range(1, count + 1)))
    row += 1
    for data in spec['rows']:
        height = 24
        for index, col in enumerate(columns, 1):
            value = _xlsx_value(data.get(col['key'], ''), col)
            cell = sheet.cell(row, index, value)
            if isinstance(value, str):
                cell.data_type = 's'
            cell.number_format = col.get('number_format', '0.00%' if col.get('type') == 'percentage' else 'm/d/yyyy' if isinstance(value, date) else 'General')
            cell.font = font
            cell.border = border
            cell.alignment = Alignment(horizontal=col.get('align', 'left'), vertical='top', wrap_text=True)
            height = max(height, fitted_height(cell, index, index))
        sheet.row_dimensions[row].height = min(409, height)
        row += 1
    last_data_row = row - 1
    if not spec['rows']:
        merge_text(row, 'No records match the selected filters.')
        row += 1
    for section in spec.get('text_sections', []):
        sheet.row_dimensions[row].height = 8
        row += 1
        merge_text(row, section['title'], bold=True, size=11)
        row += 1
        text = str(section.get('text') or 'No description recorded.')
        for part in text_parts(text):
            merge_text(row, part)
            row += 1
    sheet.auto_filter.ref = f'A{header_row}:{get_column_letter(count)}{last_data_row}'
    sheet.freeze_panes = f'A{header_row + 1}'
    sheet.print_title_rows = f'1:{header_row}'
    sheet.print_area = f'A1:{get_column_letter(count)}{row - 1}'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape' if total_width > 110 else 'portrait'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.oddFooter.center.text = 'DefenSYS | Page &P of &N'
    if include_signatures and spec.get('signatures', True):
        _write_xlsx_signatories(sheet, signatories, generated_by)
    return sheet


def rubric_document_sheets(detail):
    rubric = detail['rubric']
    criteria = detail['criteria']
    return [{
        'name': 'Rubric Details', 'title': rubric['name'], 'subtitle': 'Official Evaluation Rubric',
        'metadata': [{'label': label, 'value': rubric[key]} for label, key in (
            ('Publication Status', 'status'), ('Grading Assignment', 'assignment_status'),
            ('Academic Scope', 'scope'), ('Academic Period', 'semester'), ('Rubric Stage / Event', 'stage'),
            ('Evaluator Type', 'evaluation_type'), ('Evaluation Target', 'target'),
            ('Actual Grading Selections', 'assigned_to'), ('Assigned Grade Distribution', 'weights'))],
        'section_title': 'Evaluation Criteria',
        'caption': 'Criterion weights are configured numeric values; grade distribution is shown separately above.',
        'columns': [
            {'key': 'no', 'label': 'No.', 'width': 6}, {'key': 'criterion', 'label': 'Criterion', 'width': 30},
            {'key': 'scale', 'label': 'Scoring Scale', 'width': 18}, {'key': 'max', 'label': 'Maximum Score', 'width': 12},
            {'key': 'weight', 'label': 'Configured Weight', 'width': 14, 'type': 'number', 'number_format': '0.00'}, {'key': 'target', 'label': 'Target', 'width': 20},
        ],
        'rows': [{'no': index, 'criterion': criterion.name, 'scale': criterion.get_scale_display(),
                  'max': criterion.max_score, 'weight': criterion.weight, 'target': criterion.get_target_type_display()}
                 for index, criterion in enumerate(criteria, 1)],
        'text_sections': [{'title': f'{index}. {criterion.name} - Description', 'text': criterion.description}
                          for index, criterion in enumerate(criteria, 1)],
    }]
