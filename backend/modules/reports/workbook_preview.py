"""Print previews read from the saved XLSX, including its cells and formatting.

This renderer covers the report templates produced by DefenSYS. It needs no
desktop Office installation and never substitutes the separate PDF report.
"""
import base64
import hashlib
import io
import math
import re
from datetime import date, datetime
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.utils import range_boundaries
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A3, A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Table, TableStyle


def display_value(cell):
    value, fmt = cell.value, cell.number_format
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        if fmt.startswith('yyyy-mm-dd'):
            text = value.strftime('%Y-%m-%d')
        else:
            text = f'{value.month}/{value.day}/{value.year}'
        return text + (value.strftime(' %H:%M') if 'hh' in fmt.lower() and isinstance(value, datetime) else '')
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, (int, float)):
        decimals = re.search(r'0\.(0+)', fmt)
        if '%' in fmt:
            return f'{value * 100:.{len(decimals[1]) if decimals else 0}f}%'
        if decimals:
            return f'{value:.{len(decimals[1])}f}'
        return f'{value:g}'
    return str(value)


def font_name(font):
    serif = 'times' in (font.name or '').lower()
    if serif:
        return 'Times-BoldItalic' if font.b and font.i else 'Times-Bold' if font.b else 'Times-Italic' if font.i else 'Times-Roman'
    return 'Helvetica-BoldOblique' if font.b and font.i else 'Helvetica-Bold' if font.b else 'Helvetica-Oblique' if font.i else 'Helvetica'


def rgb_color(value, default=colors.black):
    if value is not None and value.type == 'rgb' and isinstance(value.rgb, str):
        return colors.HexColor('#' + value.rgb[-6:])
    return default


def column_points(sheet, column):
    from openpyxl.utils import get_column_letter
    dimension = sheet.column_dimensions[get_column_letter(column)]
    return 0 if dimension.hidden else (math.floor((dimension.width or 13) * 7) + 5) * .75


def _sheet_bounds(sheet):
    if sheet.print_area:
        return range_boundaries(str(sheet.print_area).split('!')[-1].split(',')[0])
    return 1, 1, sheet.max_column, sheet.max_row


def _render_sheet(sheet):
    first_col, first_row, last_col, last_row = _sheet_bounds(sheet)
    page_size = A3 if sheet.page_setup.paperSize == sheet.PAPERSIZE_A3 else A4
    if sheet.page_setup.orientation == 'landscape':
        page_size = landscape(page_size)
    margin = 24
    available_width = page_size[0] - margin * 2
    columns = [col for col in range(first_col, last_col + 1) if column_points(sheet, col)]
    total_width = sum(column_points(sheet, col) for col in columns)
    horizontal_pages = max(1, sheet.page_setup.fitToWidth or 1)
    scale = min(1, available_width * horizontal_pages / max(total_width, 1))
    chunks, chunk, used_width = [], [], 0
    for col in columns:
        width = column_points(sheet, col) * scale
        if chunk and used_width + width > available_width + .1:
            chunks.append(chunk)
            chunk, used_width = [], 0
        chunk.append(col)
        used_width += width
    if chunk:
        chunks.append(chunk)
    repeated = re.search(r'\$?(\d+):\$?(\d+)', sheet.print_title_rows or '')
    repeated_rows = int(repeated[2]) - first_row + 1 if repeated and int(repeated[1]) == first_row else 0
    merged = list(sheet.merged_cells.ranges)
    story = []
    header_rows = list(range(first_row, first_row + repeated_rows))
    available_height = page_size[1] - margin - 32 - 16
    header_height = sum((sheet.row_dimensions[row].height or 15) * scale for row in header_rows)
    groups, current, used_height = [], [], header_height
    for row in range(first_row + repeated_rows, last_row + 1):
        height = (sheet.row_dimensions[row].height or 15) * scale
        if current and used_height + height > available_height:
            groups.append(header_rows + current)
            current, used_height = [], header_height
        current.append(row)
        used_height += height
    if current:
        groups.append(header_rows + current)
    elif header_rows:
        groups.append(header_rows)
    for selected_columns in chunks:
        for selected_rows in groups:
            if story:
                story.append(PageBreak())
            story.append(_sheet_table(sheet, selected_columns, selected_rows, scale, merged))
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=page_size, leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=32,
                            title=sheet.title, author='DefenSYS')
    def footer(canvas, document):
        canvas.setFont('Helvetica', 8)
        canvas.setFillColor(colors.HexColor('#64748B'))
        canvas.drawString(margin, 15, f'DefenSYS | {sheet.title}')
        canvas.drawRightString(page_size[0] - margin, 15, f'Page {document.page}')
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def _sheet_table(sheet, selected_columns, selected_rows, scale, merged):
    # Build each physical page independently. Reusing a split merged-cell Table
    # loses its repeated headings in ReportLab's continuation-page coordinates.
    col_positions = {col: index for index, col in enumerate(selected_columns)}
    row_positions = {row: index for index, row in enumerate(selected_rows)}
    widths = [column_points(sheet, col) * scale for col in selected_columns]
    heights = [(sheet.row_dimensions[row].height or sheet.sheet_format.defaultRowHeight or 15) * scale for row in selected_rows]
    commands = [('VALIGN', (0, 0), (-1, -1), 'TOP')]
    for side in ('LEFT', 'RIGHT', 'TOP', 'BOTTOM'):
        commands.append((side + 'PADDING', (0, 0), (-1, -1), 3 * scale))
    data = []
    for row in selected_rows:
        values = []
        for col in selected_columns:
            cell = sheet.cell(row, col)
            merge = next((area for area in merged if area.min_row <= row <= area.max_row and area.min_col <= col <= area.max_col), None)
            start_col = max(merge.min_col, selected_columns[0]) if merge else col
            end_col = min(merge.max_col, selected_columns[-1]) if merge else col
            if merge and (row != merge.min_row or col != start_col):
                values.append('')
                continue
            if merge:
                cell = sheet.cell(merge.min_row, merge.min_col)
            start = (col_positions[start_col], row_positions[row])
            end = (col_positions[end_col], row_positions.get(merge.max_row, row_positions[row]) if merge else row_positions[row])
            if merge:
                commands.append(('SPAN', start, end))
            size = (cell.font.sz or 11) * scale
            alignment = cell.alignment.horizontal
            if alignment in ('general', None):
                alignment = 'right' if isinstance(cell.value, (int, float)) else 'left'
            style = ParagraphStyle('Cell', fontName=font_name(cell.font), fontSize=size, leading=size * 1.15,
                alignment={'center': TA_CENTER, 'right': TA_RIGHT}.get(alignment, TA_LEFT),
                textColor=rgb_color(cell.font.color), splitLongWords=True)
            text = escape(display_value(cell)).replace('\n', '<br/>')
            values.append(Paragraph(text or '&#160;', style))
            commands.append(('VALIGN', start, end, {'center': 'MIDDLE', 'bottom': 'BOTTOM'}.get(cell.alignment.vertical, 'TOP')))
            if cell.fill.patternType == 'solid':
                commands.append(('BACKGROUND', start, end, rgb_color(cell.fill.fgColor, colors.white)))
            for side, command in (('left', 'LINEBEFORE'), ('right', 'LINEAFTER'), ('top', 'LINEABOVE'), ('bottom', 'LINEBELOW')):
                border = getattr(cell.border, side)
                if border and border.style:
                    commands.append((command, start, end, (.8 if border.style in ('medium', 'thick') else .4) * scale,
                                     rgb_color(border.color)))
        data.append(values)
    table = Table(data, colWidths=widths, rowHeights=heights, hAlign='LEFT')
    table.setStyle(TableStyle(commands))
    return table


def build_workbook_preview(xlsx_bytes):
    """Reopen the exact XLSX bytes; preview every visible worksheet's print area."""
    workbook = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    return {
        'sha256': hashlib.sha256(xlsx_bytes).hexdigest(),
        'sheets': [{'name': sheet.title, 'pdf_base64': base64.b64encode(_render_sheet(sheet)).decode('ascii')}
                   for sheet in workbook if sheet.sheet_state == 'visible'],
    }
