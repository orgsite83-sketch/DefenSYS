import io
from datetime import date, time
from pathlib import Path
from unittest.mock import patch

import openpyxl
from django.contrib.auth import get_user_model
from django.utils import timezone
from rest_framework.test import APITestCase

from academic_period_management.models import SchoolYear, Semester
from student_teams.models import StudentTeam, TeamMembership
from grading.rubrics.models import Rubric, RubricCriterion
from defense.stages.models import DefenseStage, StageGradingConfig, StageDeliverable
from defense.scheduler.models import DefenseSchedule, PitEventGradingConfig
from user_management.academic_records.models import StudentAcademicRecord
from user_management.models import ExternalEvaluator, GuestPanelistCode
from repository.archive.models import ArchiveEntry
from reports.catalog_reports import BUILDERS, render_pdf, column, dataset


class CatalogReportTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        User = get_user_model()
        cls.admin = User.objects.create_user(username='export-admin', role='admin', first_name='Export', last_name='Administrator')
        cls.lead = User.objects.create_user(username='export-lead', role='faculty', is_pit_lead=True, pit_lead_year='2nd Year')
        cls.student = User.objects.create_user(username='2026-export-student', role='student', first_name='Andrea', last_name='Santos')
        cls.year = SchoolYear.objects.create(label='2040-2041')
        cls.sem = Semester.objects.create(school_year=cls.year, label='1st Semester', is_active=True)
        cls.old_sem = Semester.objects.create(school_year=cls.year, label='2nd Semester')
        cls.stage, _ = DefenseStage.objects.get_or_create(label='Export Proposal', defaults={'display_order': 1})
        cls.team = StudentTeam.objects.create(name='Export Team', project_title='Patient Records System',
            level=StudentTeam.LEVEL_4_CAPSTONE, year_level='4th Year', section='BSIT-4A', semester=cls.sem, leader=cls.student)
        cls.pit_team = StudentTeam.objects.create(name='PIT Export Team', project_title='Smart Campus',
            level=StudentTeam.LEVEL_2_PIT, year_level='2nd Year', section='BSIT-2A', semester=cls.sem, leader=cls.student)
        cls.other_pit = StudentTeam.objects.create(name='Other Year Team', project_title='Other project',
            level=StudentTeam.LEVEL_1_PIT, year_level='1st Year', semester=cls.sem, leader=cls.student)
        TeamMembership.objects.create(team=cls.team, student=cls.student)
        cls.rubric = Rubric.objects.create(name='Proposal Quality Guide', scope='capstone', semester=cls.sem,
            defense_stage=cls.stage, evaluation_type='panel', status='published', created_by=cls.admin)
        cls.description = ('Design evidence <verified> & documented. ' * 140) + 'COMPLETE DESCRIPTION END'
        RubricCriterion.objects.create(rubric=cls.rubric, name='System Design', description=cls.description,
            scale=Rubric.SCALE_10, max_score=10, weight=100, display_order=1)
        StageGradingConfig.objects.create(defense_stage=cls.stage, semester=cls.sem, panel_rubric=cls.rubric,
            panel_weight=60, adviser_weight=25, peer_weight=15)
        StageDeliverable.objects.create(defense_stage=cls.stage, deliverable_id='D1', label='Proposal Paper', required=True)
        cls.pit_rubric = Rubric.objects.create(name='IT Olympics Scoring', scope='pit', semester=cls.sem,
            event_name='IT Olympics', evaluation_type='panel', status='published', created_by=cls.lead)
        PitEventGradingConfig.objects.create(year_level='2nd Year', semester=cls.sem, event_name='IT Olympics', panel_rubric=cls.pit_rubric,
            panel_weight=75, peer_weight=25)
        StudentAcademicRecord.objects.create(student=cls.student, semester=cls.sem, year_level='4th Year', section='BSIT-4A')
        cls.archive = ArchiveEntry.objects.create(entry_type='capstone', team=cls.team,
            file_name='Student Performance Analytics and Early Warning with Complete Filename.pdf',
            team_name=cls.team.name, year_level='4th Year', academic_year=cls.year.label,
            semester_label=cls.sem.label, stage_label=cls.stage.label, uploaded_by=cls.admin, uploaded_at=timezone.now())
        cls.schedule = DefenseSchedule.objects.create(team=cls.team, semester=cls.sem, defense_stage=cls.stage,
            scheduled_date=date(2040, 10, 15), start_time=time(9, 30), room='Room A', scope='capstone')
        cls.evaluator = ExternalEvaluator.objects.create(name='Dr. Maria Rivera', email='export@example.test', institution='Partner University', status='approved', created_by=cls.admin)
        cls.invitation = GuestPanelistCode.objects.create(code='SECRET-EXPORT', evaluator=cls.evaluator,
            guest_name=cls.evaluator.name, defense_schedule=cls.schedule, created_by=cls.admin)

    def setUp(self):
        self.client.force_authenticate(self.admin)

    def report(self, path, params=None):
        return self.client.get(path, {'export_format': 'json', **(params or {})})

    def test_every_catalog_report_supports_preview_pdf_and_excel(self):
        for kind in BUILDERS:
            for format_value in ('json', 'pdf', 'xlsx'):
                with self.subTest(report=kind, format=format_value):
                    response = self.report(f'/api/reports/{kind}/', {'export_format': format_value})
                    self.assertEqual(response.status_code, 200, getattr(response, 'data', None))
                    if format_value == 'json':
                        self.assertTrue(response.data.get('pdf_base64'), f'{kind}: PDF preview failed')
                        self.assertIn('rows', response.data)
                        self.assertNotIn('excel_preview_error', response.data)
                        self.assertTrue(response.data['excel_workbook_preview']['sheets'])
                    elif format_value == 'pdf':
                        self.assertTrue(response.content.startswith(b'%PDF'))
                        self.assert_official_template(response.content)
                        output = Path(__file__).resolve().parents[3] / '.tmp'
                        output.mkdir(exist_ok=True)
                        (output / f'{kind}-official-template-qa.pdf').write_bytes(response.content)
                    else:
                        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
                        for sheet in workbook:
                            if kind in ('defense-stage-setup', 'pit-event-setup'):
                                header_row = int(sheet.auto_filter.ref.split(':')[0][1:])
                                self.assertEqual(sheet.cell(header_row, 1).value, 'Deliverable Code')
                                self.assertEqual(sheet.freeze_panes, f'A{header_row + 1}')
                            else:
                                self.assertEqual(sheet.freeze_panes, 'A4')
                                self.assertTrue(sheet.auto_filter.ref.startswith('A3:'))
                            self.assertTrue(sheet.sheet_view.showGridLines)
                            self.assertTrue(sheet['A1'].value)
                            self.assertTrue(sheet['A2'].value)

    def assert_official_template(self, content):
        import pdfplumber
        with pdfplumber.open(io.BytesIO(content)) as document:
            self.assertTrue(document.pages)
            for page in document.pages:
                self.assertEqual((page.width, page.height), (612, 792))
                # The full-width institutional letterhead and left policy panel
                # must be present on every page, just as in the grade reports.
                self.assertTrue(any(abs(image['x0']) < 1 and abs(image['width'] - 612) < 1 for image in page.images))
                self.assertTrue(any(abs(image['x0'] - 24) < 1 and abs(image['width'] - 93.6) < 1 for image in page.images))
                self.assertIn('Page ', page.extract_text() or '')

    def test_wide_pdf_register_keeps_every_field_readable_across_pages(self):
        columns = [column(f'field_{index}', f'Column {index}', 24, type='datetime' if index == 11 else 'text')
                   for index in range(12)]
        rows = [{f'field_{index}': ('2040-10-15T01:30:00+00:00' if index == 11 else f'Record {record} / Value {index}')
                 for index in range(12)} for record in range(1, 13)]
        spec = dataset('Wide Register', columns, rows)
        content = render_pdf(spec, [{'label': 'Records', 'value': 12}], 'Export Administrator', False, [])
        output = Path(__file__).resolve().parents[3] / '.tmp'
        output.mkdir(exist_ok=True)
        (output / 'wide-register-official-template-qa.pdf').write_bytes(content)
        self.assert_official_template(content)
        import pdfplumber
        with pdfplumber.open(io.BytesIO(content)) as document:
            text = '\n'.join(page.extract_text() or '' for page in document.pages)
            for index in range(12):
                self.assertIn(f'Column {index}', text)
            self.assertIn('Record 12 / Value 10', text)
            self.assertIn('09:30 AM', text)
            self.assertNotIn('+00:00', text)
            self.assertGreater(len(document.pages), 1)
            for page in document.pages:
                self.assertTrue(all(0 <= char['x0'] < char['x1'] <= 612 for char in page.chars))
        empty = render_pdf(dataset('Empty Register', columns, []), [{'label': 'Records', 'value': 0}], 'Export Administrator', False, [])
        self.assert_official_template(empty)

    def test_rubric_export_preserves_complete_criteria_and_effective_weights(self):
        response = self.report(f'/api/reports/rubrics/{self.rubric.pk}/', {'export_format': 'xlsx'})
        self.assertEqual(response.status_code, 200)
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ['Rubric Details'])
        main_text = str(list(workbook.active.values))
        self.assertIn('Panel: 60%', main_text)
        sheet = workbook.active
        description_start = next(cell.row for row in sheet for cell in row if cell.value == '1. System Design - Description')
        description = ''.join(str(sheet.cell(row, 1).value or '') for row in range(description_start + 1, sheet.max_row + 1))
        self.assertEqual(description, self.description)
        criterion_row = next(row for row in sheet if row[1].value == 'System Design')
        self.assertEqual(criterion_row[4].value, 100)
        self.assertEqual(criterion_row[4].number_format, '0.00')
        preview = self.report(f'/api/reports/rubrics/{self.rubric.pk}/').data
        self.assertNotIn('excel_preview_error', preview)
        self.assertEqual(preview['excel_workbook_preview']['sheets'][0]['name'], 'Rubric Details')
        pdf = self.report(f'/api/reports/rubrics/{self.rubric.pk}/', {'export_format': 'pdf'})
        self.assertTrue(pdf.content.startswith(b'%PDF'))
        output = Path(__file__).resolve().parents[3] / '.tmp'
        output.mkdir(exist_ok=True)
        (output / 'rubric-export-qa.pdf').write_bytes(pdf.content)
        import base64
        import pdfplumber
        with pdfplumber.open(io.BytesIO(base64.b64decode(preview['excel_workbook_preview']['sheets'][0]['pdf_base64']))) as document:
            self.assertTrue(all(char['top'] >= 20 for page in document.pages for char in page.chars))
            self.assertTrue(all(self.rubric.name in (page.extract_text() or '') for page in document.pages))
            self.assertIn('COMPLETE DESCRIPTION END', '\n'.join(page.extract_text() or '' for page in document.pages))
        (output / 'rubric-workbook-qa.xlsx').write_bytes(response.content)
        (output / 'rubric-workbook-preview-qa.pdf').write_bytes(base64.b64decode(preview['excel_workbook_preview']['sheets'][0]['pdf_base64']))

    def test_rubric_publication_is_separate_from_actual_grading_assignment_and_exports_are_read_only(self):
        draft = Rubric.objects.create(name='Draft Proposal Guide', scope='capstone', semester=self.sem,
            defense_stage=self.stage, evaluation_type='adviser', status='draft', created_by=self.admin)
        published = Rubric.objects.create(name='Published Unassigned Guide', scope='capstone', semester=self.sem,
            defense_stage=self.stage, evaluation_type='peer', status='published', created_by=self.admin)
        unused_stage = DefenseStage.objects.create(label='Unconfigured Export Stage', display_order=99)
        Rubric.objects.create(name='Unconfigured Stage Guide', scope='capstone', semester=self.sem,
            defense_stage=unused_stage, evaluation_type='panel', status='draft', created_by=self.admin)
        count = StageGradingConfig.objects.count()
        with patch('defense.stages.grading_config.get_or_create_stage_grading_config', side_effect=AssertionError('Exports must not create grading configuration')):
            preview = self.report('/api/reports/rubrics/', {'scope': 'capstone'}).data
            rows = {row['name']: row for row in preview['rows']}
            self.assertEqual(rows[draft.name]['status'], 'Draft')
            self.assertEqual(rows[draft.name]['assignment_status'], 'Not assigned')
            self.assertEqual(rows[published.name]['status'], 'Published')
            self.assertEqual(rows[published.name]['assignment_status'], 'Not assigned')
            self.assertEqual(rows[published.name]['stage'], self.stage.label)
            self.assertEqual(rows[published.name]['weights'], 'No grading assignment')
            self.assertEqual(rows[self.rubric.name]['assignment_status'], 'Assigned')
            self.assertIn('Panel: 60%', rows[self.rubric.name]['weights'])
            for fmt in ('pdf', 'xlsx'):
                response = self.report('/api/reports/rubrics/', {'scope': 'capstone', 'export_format': fmt})
                self.assertEqual(response.status_code, 200)
                output = Path(__file__).resolve().parents[3] / '.tmp'
                output.mkdir(exist_ok=True)
                (output / f'rubric-status-catalog-qa.{fmt}').write_bytes(response.content)
        self.assertEqual(StageGradingConfig.objects.count(), count)
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        main = workbook['Report Data']
        headers = [cell.value for cell in main[3]]
        self.assertIn('Publication Status', headers)
        self.assertIn('Grading Assignment', headers)
        self.assertEqual(workbook.sheetnames, ['Report Data', 'Assignments', 'Criteria'])
        assigned_row = next(row for row in main if row[1].value == self.rubric.name)
        self.assertTrue(assigned_row[headers.index('Grading Assignment')].hyperlink.location.startswith("'Assignments'!"))
        self.assertTrue(assigned_row[1].hyperlink.location.startswith("'Criteria'!"))
        StageGradingConfig.objects.filter(defense_stage=self.stage, semester=self.sem).update(adviser_rubric=draft)
        row = next(row for row in self.report('/api/reports/rubrics/').data['rows'] if row['name'] == draft.name)
        self.assertEqual((row['status'], row['assignment_status']), ('Draft', 'Assigned'))

    def test_rubric_export_lists_all_actual_stage_and_event_assignments(self):
        StageGradingConfig.objects.create(defense_stage=self.stage, semester=self.old_sem,
            panel_rubric=self.rubric, panel_weight=50, adviser_weight=30, peer_weight=20)
        PitEventGradingConfig.objects.create(year_level='2nd Year', semester=self.sem, event_name='IT Hackathon',
            panel_rubric=self.pit_rubric, panel_weight=60, peer_weight=40)
        preview = self.report('/api/reports/rubrics/', {'term_context': 'all'}).data
        rows = {row['name']: row for row in preview['rows']}
        self.assertIn(self.old_sem.display_name, rows[self.rubric.name]['assigned_to'])
        self.assertIn('IT Olympics', rows[self.pit_rubric.name]['assigned_to'])
        self.assertIn('IT Hackathon', rows[self.pit_rubric.name]['assigned_to'])
        self.assertEqual(len([row for row in preview['sections'][0]['rows'] if row['rubric'] == self.rubric.name]), 2)
        pdf = self.report('/api/reports/rubrics/', {'export_format': 'pdf'})
        import pdfplumber
        with pdfplumber.open(io.BytesIO(pdf.content)) as document:
            text = '\n'.join(page.extract_text() or '' for page in document.pages)
            self.assertIn('Rubric Inventory', text)
            self.assertIn('Publication Status', text)
            self.assertIn('Actual Grading Assignments', text)
            self.assertIn('IT Hackathon', text)
            # The inventory page precedes the groups and must not inherit
            # the last rubric's continuation footer during canvas replay.
            self.assertFalse(any('Rubric ' in line and ' - ' in line
                                 for line in (document.pages[0].extract_text() or '').splitlines()))
            self.assertTrue(all('Proposal Quality Guide' in (page.extract_text() or '') for page in document.pages
                                if 'Design evidence' in (page.extract_text() or '')))

    def test_stage_setup_groups_each_stage_with_its_rubrics_and_requirements(self):
        second = DefenseStage.objects.create(label='Export Final Defense', display_order=98)
        StageDeliverable.objects.create(defense_stage=second, deliverable_id='F1', label='Final Manuscript', required=True)
        params = {'semester_id': self.sem.pk, 'search': 'Export'}
        pdf = self.report('/api/reports/defense-stage-setup/', {**params, 'export_format': 'pdf'})
        import pdfplumber
        with pdfplumber.open(io.BytesIO(pdf.content)) as document:
            proposal = next(page.extract_text() or '' for page in document.pages if 'Proposal Paper' in (page.extract_text() or ''))
            final = next(page.extract_text() or '' for page in document.pages if 'Final Manuscript' in (page.extract_text() or ''))
            self.assertIn(self.stage.label, proposal)
            self.assertIn(self.rubric.name, proposal)
            self.assertNotIn('Final Manuscript', proposal)
            self.assertIn(second.label, final)
            self.assertNotIn('Proposal Paper', final)
            self.assertNotIn('Record numbers link', proposal)
        xlsx = self.report('/api/reports/defense-stage-setup/', {**params, 'export_format': 'xlsx'})
        output = Path(__file__).resolve().parents[3] / '.tmp'
        output.mkdir(exist_ok=True)
        (output / 'grouped-stage-config-qa.pdf').write_bytes(pdf.content)
        (output / 'grouped-stage-config-qa.xlsx').write_bytes(xlsx.content)
        import base64
        preview = self.report('/api/reports/defense-stage-setup/', params).data
        for index, sheet in enumerate(preview['excel_workbook_preview']['sheets'], 1):
            (output / f'grouped-stage-workbook-preview-{index}-qa.pdf').write_bytes(base64.b64decode(sheet['pdf_base64']))
        workbook = openpyxl.load_workbook(io.BytesIO(xlsx.content))
        self.assertEqual(workbook.sheetnames, [self.stage.label, second.label])
        self.assertIn('Proposal Paper', str(list(workbook[self.stage.label].values)))
        self.assertNotIn('Final Manuscript', str(list(workbook[self.stage.label].values)))
        self.assertIn('Final Manuscript', str(list(workbook[second.label].values)))

    def test_stage_dropdown_options_supply_names_scopes_and_periods(self):
        response = self.report('/api/reports/catalog-options/', {'report': 'defense-stage-setup'})
        self.assertIn({'value': self.stage.label, 'label': self.stage.label, 'scope': 'capstone'}, response.data['stage_choices'])
        self.assertIn({'value': 'IT Olympics', 'label': 'IT Olympics', 'scope': 'pit', 'semester_id': self.sem.pk}, response.data['stage_choices'])

    def test_project_register_filters_all_rows_and_uses_excel_dates(self):
        params = {'search': 'Student Performance', 'type': 'capstone', 'year_level': '4th Year', 'academic_year': self.year.label, 'export_format': 'xlsx', 'page_size': 1}
        response = self.report('/api/reports/project-records/', params)
        self.assertEqual(response.status_code, 200)
        worksheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
        rows = [row for row in worksheet.iter_rows() if row[1].value == self.archive.file_name]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][-1].data_type, 'd')
        self.assertGreaterEqual(worksheet.column_dimensions['B'].width, 50)
        self.assertTrue(rows[0][1].alignment.wrap_text)

    def test_enrollment_local_filters_match_source_records(self):
        response = self.report('/api/reports/student-enrollment/', {'section': 'BSIT-4A', 'year_level': '4th Year'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual([row['name'] for row in response.data['rows']], ['Andrea Santos'])
        empty = self.report('/api/reports/student-enrollment/', {'section': 'UNASSIGNED'})
        self.assertEqual(empty.data['rows'], [])
        self.assertTrue(empty.data['pdf_base64'])

    def test_readiness_does_not_mutate_schedules_or_teams(self):
        with patch.object(DefenseSchedule, 'save', side_effect=AssertionError('Export must not normalize schedules')), patch.object(StudentTeam, 'save', side_effect=AssertionError('Export must not modify teams')):
            response = self.report('/api/reports/readiness/', {'stage': self.stage.label, 'scope': 'capstone'})
            self.assertEqual(response.status_code, 200)
            self.assertEqual([row['team'] for row in response.data['rows']], [self.team.name])

    def test_pit_lead_cannot_widen_year_or_capstone_visibility(self):
        self.client.force_authenticate(self.lead)
        response = self.report('/api/reports/overview/', {'scope': 'capstone', 'year_level': '4th Year'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['rows'][0]['value'], 0)
        options = self.report('/api/reports/catalog-options/')
        self.assertEqual(options.status_code, 200)
        self.assertEqual([row['id'] for row in options.data['teams']], [self.pit_team.pk])
        self.assertNotIn(self.rubric.pk, [row['id'] for row in options.data['rubrics']])
        denied = self.report('/api/reports/academic-periods/')
        self.assertEqual(denied.status_code, 403)
        rubric = self.report(f'/api/reports/rubrics/{self.rubric.pk}/')
        self.assertEqual(rubric.status_code, 404)

    def test_external_assignments_have_names_dates_and_no_credentials(self):
        response = self.report('/api/reports/external-evaluators/')
        self.assertEqual(response.status_code, 200)
        self.assertNotIn(self.invitation.code, str(response.data))
        row = response.data['sections'][0]['rows'][0]
        self.assertEqual(row['date'], '2040-10-15')
        self.assertEqual(row['time'], '09:30')
        self.assertEqual(row['team'], self.team.name)

    def test_invalid_ids_and_missing_selections_are_controlled(self):
        for kind in BUILDERS:
            response = self.report(f'/api/reports/{kind}/', {'semester_id': 'broken'})
            self.assertEqual(response.status_code, 400)
        self.assertEqual(self.report('/api/reports/rubrics/999999/').status_code, 404)
        for kind in ('semester-grades', 'team-roster', 'defense-schedules'):
            self.assertEqual(self.report(f'/api/reports/{kind}/', {'semester_id': 'broken'}).status_code, 400)

    def test_source_roster_and_board_filters_are_applied(self):
        roster = self.report('/api/reports/team-roster/', {'search': 'Patient', 'section': 'BSIT-4A'})
        self.assertEqual([row['team_name'] for row in roster.data['rows']], [self.team.name])
        schedule = self.report('/api/reports/defense-schedules/', {'source_context': 'defense-board', 'section': 'BSIT-2A'})
        self.assertEqual(schedule.data['rows'], [])

    def test_nonmanager_cannot_export_private_registries(self):
        self.client.force_authenticate(self.student)
        for kind in ('academic-periods', 'faculty-directory', 'student-enrollment', 'project-records', 'external-evaluators', 'readiness', 'overview', 'curriculum-summary'):
            self.assertEqual(self.report(f'/api/reports/{kind}/').status_code, 403)
        self.client.force_authenticate(None)
        self.assertEqual(self.report('/api/reports/rubrics/').status_code, 401)

    def test_assignment_tab_exports_only_matching_defenses(self):
        response = self.report('/api/reports/external-evaluators/', {'view': 'assignments', 'search': 'Maria'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['rows'][0]['team'], self.team.name)
        self.assertNotIn(self.invitation.code, str(response.data))
        empty = self.report('/api/reports/external-evaluators/', {'view': 'assignments', 'search': 'unmatched'})
        self.assertEqual(empty.data['rows'], [])

    def test_rubric_catalog_history_and_pit_weights(self):
        response = self.report('/api/reports/rubrics/', {'scope': 'pit'})
        self.assertEqual(response.data['rows'][0]['weights'], 'Panel: 75%, Peer: 25%')
        history = self.report('/api/reports/rubrics/', {'term_context': 'history'})
        self.assertEqual(history.data['rows'], [])

    def test_all_history_roster_export_retains_period_coverage(self):
        historical = StudentTeam.objects.create(name='Historical Export Team', project_title='Old Project',
            level=StudentTeam.LEVEL_4_CAPSTONE, year_level='4th Year', semester=self.old_sem, leader=self.student)
        response = self.report('/api/reports/team-roster/', {'term_context': 'all', 'search': 'Historical'})
        self.assertEqual([row['team_name'] for row in response.data['rows']], [historical.name])
