"""Workbook behaviour that affects editing, filtering and readability in Excel."""
import io
from datetime import datetime

import openpyxl
from django.test import SimpleTestCase

from reports.export_formatters import build_preview_response, generate_xlsx_response
from reports.workbook_preview import build_workbook_preview


class ExcelLayoutTests(SimpleTestCase):
    metadata = [
        {'label': 'Academic Period', 'value': '1st Semester, A.Y. 2026-2027'},
        {'label': 'Scope', 'value': 'capstone'},
        {'label': 'Stage', 'value': 'Concept Proposal'},
        {'label': 'Records', 'value': 2},
    ]
    columns = [
        {'key': 'student', 'label': 'Student'},
        {'key': 'score', 'label': 'Final Grade', 'type': 'percentage'},
        {'key': 'count', 'label': 'Count', 'type': 'integer'},
        {'key': 'date', 'label': 'Date', 'type': 'date'},
        {'key': 'uploaded', 'label': 'Uploaded At', 'type': 'datetime'},
        {'key': 'notes', 'label': 'Notes', 'width': 30},
    ]
    rows = [
        {'student': '00123', 'score': '83.75%', 'count': '4', 'date': '2026-06-18',
         'uploaded': '2026-06-18T01:00:00Z', 'notes': '=HYPERLINK("unsafe")'},
        {'student': 'Student Two', 'score': 'Pending', 'count': '0', 'date': '',
         'uploaded': '', 'notes': 'Long description ' * 50},
    ]

    def workbook(self, **kwargs):
        response = generate_xlsx_response('Capstone Grades', self.metadata, self.columns, self.rows, 'grades', **kwargs)
        return openpyxl.load_workbook(io.BytesIO(response.content))

    def test_compact_merged_context_and_editable_filter_table(self):
        sheet = self.workbook().active
        self.assertEqual(sheet['A1'].value, 'Capstone Grades')
        self.assertEqual({str(area) for area in sheet.merged_cells.ranges}, {'A1:F1', 'A2:F2'})
        self.assertIn('Concept Proposal', sheet['A2'].value)
        self.assertNotIn('Records:', sheet['A2'].value)
        self.assertEqual(sheet.auto_filter.ref, 'A3:F5')
        self.assertEqual(sheet.freeze_panes, 'A4')
        self.assertTrue(sheet.sheet_view.showGridLines)
        self.assertFalse(sheet.protection.sheet)
        self.assertEqual(sheet.sheet_view.zoomScale, 100)
        self.assertEqual(sheet.print_title_rows, '$1:$3')
        self.assertIn('$A$1:$F$5', sheet.print_area)
        for row in sheet.iter_rows(min_row=3):
            for cell in row:
                self.assertEqual(cell.border.left.style, 'thin')
                self.assertEqual(cell.border.left.color.rgb, '00000000')
                self.assertIsNone(cell.fill.patternType)
        self.assertTrue(all(cell.font.bold for cell in sheet[3]))
        self.assertGreater(sheet.row_dimensions[5].height, sheet.row_dimensions[4].height)
        self.assertTrue(sheet['F5'].alignment.wrap_text)
        self.assertIn('Generated', sheet['A1'].comment.text)

    def test_numeric_dates_pending_and_literal_identifiers(self):
        sheet = self.workbook().active
        self.assertEqual(sheet['A4'].value, '00123')
        self.assertEqual(sheet['B4'].value, 0.8375)
        self.assertEqual(sheet['B4'].number_format, '0.00%')
        self.assertEqual(sheet['B5'].value, 'Pending')
        self.assertEqual(sheet['C4'].value, 4)
        self.assertEqual(sheet['C5'].value, 0)
        self.assertEqual(sheet['D4'].value, datetime(2026, 6, 18))
        self.assertEqual(sheet['D4'].number_format, 'm/d/yyyy')
        self.assertEqual(sheet['E4'].value, datetime(2026, 6, 18, 9))
        self.assertEqual(sheet['F4'].data_type, 's')
        self.assertFalse(any(cell.data_type == 'f' for row in sheet for cell in row))

    def test_supporting_sheets_and_preview_preserve_filter_context(self):
        extras = [{'name': 'Criteria', 'title': 'Rubric Criteria', 'metadata': [],
                   'columns': self.columns, 'rows': self.rows}]
        workbook = self.workbook(extra_sheets=extras)
        self.assertEqual(workbook['Criteria']['A2'].value, workbook.active['A2'].value)
        self.assertEqual(workbook['Criteria'].auto_filter.ref, 'A3:F5')
        payload = build_preview_response('Capstone Grades', '', [], self.metadata, self.columns, self.rows, extra_sheets=extras).data
        self.assertEqual(payload['excel_header'], [workbook.active['A1'].value, workbook.active['A2'].value])
        self.assertEqual(payload['excel_sheets'][0]['excel_header'], ['Rubric Criteria', workbook['Criteria']['A2'].value])

    def test_zero_percent_and_empty_workbook(self):
        response = generate_xlsx_response('Empty', [], self.columns, [], 'empty')
        sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual(sheet.auto_filter.ref, 'A3:F3')
        self.assertEqual(sheet['A4'].value, 'No records match the selected filters.')
        response = generate_xlsx_response('Zero', [], self.columns, [{'score': '0%'}], 'zero')
        sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(sheet['B4'].value, 0)
        self.assertEqual(sheet['B4'].data_type, 'n')

    def test_preview_reads_saved_workbook_cells_and_formats(self):
        import base64
        import hashlib
        import pdfplumber
        workbook = self.workbook()
        # Only the workbook changes; the original report JSON still says 83.75%.
        workbook.active['B4'] = 0.425
        workbook.active['B4'].number_format = '0.00%'
        workbook.active['F4'] = '=literal cell text'
        workbook.active['F4'].data_type = 's'
        output = io.BytesIO()
        workbook.save(output)
        preview = build_workbook_preview(output.getvalue())
        self.assertEqual(preview['sha256'], hashlib.sha256(output.getvalue()).hexdigest())
        with pdfplumber.open(io.BytesIO(base64.b64decode(preview['sheets'][0]['pdf_base64']))) as pdf:
            text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
        self.assertIn('42.50%', text)
        self.assertNotIn('83.75%', text)
        self.assertIn('00123', text)
        self.assertIn('Pending', text)
        self.assertIn('=literal cell text', text)
        self.assertIn('6/18/2026', text)
