import io
import json
from datetime import datetime
from unittest import skipUnless

import openpyxl
from django.contrib.auth import get_user_model
from django.utils import timezone
from rest_framework.test import APITestCase

from authentication_access_control.audit import log_high_impact_action
from authentication_access_control.models import SystemAuditLog, SystemAuditReview
from reports.export_formatters import HAS_OPENPYXL


@skipUnless(HAS_OPENPYXL, 'Excel exports require openpyxl')
class AuditExportTests(APITestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_user(username='export-auditor', role='admin')
        self.client.force_authenticate(self.admin)

    def log(self, action='rubric.create', **kwargs):
        return log_high_impact_action(
            category=SystemAuditLog.CATEGORY_RUBRICS, action=action,
            target=self.admin, target_type='Rubric', actor=self.admin,
            new_values=kwargs.pop('new_values', {'name': 'Panel rubric'}), **kwargs,
        )

    def export(self, **params):
        response = self.client.get('/api/reports/audit-trail/', {'export_format': 'xlsx', **params})
        self.assertEqual(response.status_code, 200)
        return openpyxl.load_workbook(io.BytesIO(response.content))

    @staticmethod
    def data(ws):
        header = int(ws.auto_filter.ref.split(':')[0][1:])
        last_row = openpyxl.utils.range_boundaries(ws.auto_filter.ref)[3]
        labels = [cell.value for cell in ws[header]]
        return header, [dict(zip(labels, [cell.value for cell in row])) for row in ws.iter_rows(min_row=header + 1, max_row=last_row)]

    def test_excel_signatories_follow_settings_without_changing_records(self):
        self.log()
        self.log('rubric.update', old_values={'name': 'Previous'})
        signers = [
            {'label': 'Verified by:', 'name': 'Audit Reviewer', 'role': 'Quality Officer'},
            {'label': 'Approved by:', 'name': 'Department Head', 'role': 'Program Chair'},
        ]
        unsigned = self.export(signatories=json.dumps(signers))
        signed = self.export(include_signatures='true', signatories=json.dumps(signers))
        register = signed['Audit Register']
        self.assertEqual(self.data(unsigned.active), self.data(register))
        self.assertEqual(register.auto_filter.ref, 'A3:L5')
        self.assertEqual(register.max_row, 16)
        self.assertEqual(register['A8'].value, 'Verified by:')
        self.assertEqual(register['A10'].value, 'Audit Reviewer')
        self.assertEqual(register['A11'].value, 'Quality Officer')
        self.assertEqual(register['A15'].value, 'Department Head')
        self.assertEqual(register['A16'].value, 'Program Chair')
        self.assertIn('$L$16', register.print_area)
        self.assertEqual(register.sheet_view.zoomScale, 100)
        self.assertLessEqual(max(dim.height for row, dim in register.row_dimensions.items() if row >= 8), 45)
        self.assertTrue(all(area.max_row < 3 or area.min_row > 5 for area in register.merged_cells.ranges))
        self.assertTrue(register['A10'].font.bold)
        self.assertEqual(register['A10'].border.top.style, 'thin')
        self.assertEqual(unsigned.active.max_row, 5)
        for name in ('Event Evidence', 'Review History'):
            self.assertEqual(self.data(unsigned[name]), self.data(signed[name]))
        explicitly_off = self.export(include_signatures='false', signatories=json.dumps(signers))
        self.assertEqual(explicitly_off.active.max_row, 5)

    def test_signing_empty_export_is_bounded_and_signer_text_is_literal(self):
        signers = [None, {'label': 'Signed by:', 'name': '=1+1', 'role': 'Reviewer'}]
        register = self.export(include_signatures='true', signatories=json.dumps(signers)).active
        self.assertEqual(register['A4'].value, 'No records match the selected filters.')
        self.assertEqual(register['A9'].value, '=1+1')
        self.assertEqual(register['A9'].data_type, 's')
        self.assertEqual(register.max_row, 10)
        self.assertIn('$L$10', register.print_area)

    def test_enabled_signatures_have_default_signers_when_none_are_supplied(self):
        self.log()
        values = [cell.value for row in self.export(include_signatures='true').active for cell in row]
        self.assertIn('Prepared by:', values)
        self.assertIn('export-auditor', values)
        self.assertIn('Noted by:', values)
        self.assertIn('Approved by:', values)

    def test_creation_capture_and_human_review_are_separate(self):
        log = self.log()
        self.assertEqual(log.review_status, 'captured')
        register = self.export()['Audit Register']
        header, rows = self.data(register)
        self.assertEqual(rows[0]['Evidence Status'], 'Evidence captured')
        self.assertEqual(rows[0]['Reason for Change'], 'Not required for this event')
        self.assertEqual(rows[0]['Review Status'], 'Pending review')
        self.assertEqual(rows[0]['Reviewed By'], 'Pending review')
        self.assertEqual(rows[0]['Last Review Timestamp'], 'Not reviewed')
        self.assertEqual(register['A3'].value, 'Timestamp')
        self.assertEqual(header, 3)
        self.assertNotIn('Event ID', rows[0])
        self.assertIsInstance(rows[0]['Timestamp'], datetime)
        self.assertEqual(register['A4'].number_format, 'm/d/yyyy hh:mm')
        self.assertEqual(rows[0]['Timestamp'].strftime('%Y-%m-%d %H:%M'), timezone.localtime(log.created_at).strftime('%Y-%m-%d %H:%M'))
        self.assertEqual(register.freeze_panes, f'A{header + 1}')
        self.assertTrue(register.cell(header + 1, 8).alignment.wrap_text)
        self.assertTrue(register.sheet_view.showGridLines)
        self.assertTrue(register.cell(header, 1).font.bold)
        self.assertEqual(register.cell(header, 1).font.color.rgb, '00000000')
        self.assertTrue(all(cell.fill.patternType is None for ws in register.parent for row in ws for cell in row))
        self.assertTrue(all(ws.sheet_properties.tabColor is None for ws in register.parent))
        self.assertEqual(register.print_title_rows, f'$1:${header}')
        self.assertTrue(all(area.max_row < header for area in register.merged_cells.ranges))
        self.assertEqual(register['A4'].border.left.color.rgb, '00000000')
        evidence_column = next(cell for cell in register[header] if cell.value == 'Evidence Status').column
        link = register.cell(header + 1, evidence_column).hyperlink
        self.assertIsNotNone(link)
        evidence_header, evidence_rows = self.data(register.parent['Event Evidence'])
        self.assertEqual(link.location, f"'Event Evidence'!A{evidence_header + 1}")
        self.assertEqual(evidence_rows[0]['Target Entity'], rows[0]['Target Entity'])
        self.assertNotIn('Event ID', evidence_rows[0])

        for export_format in ('json', 'csv'):
            response = self.client.get('/api/reports/audit-trail/', {'export_format': export_format})
            self.assertEqual(response.status_code, 200)
            if export_format == 'json':
                self.assertEqual(response.data['columns'][0]['key'], 'created_at')
                self.assertNotIn('event_id', [column['key'] for column in response.data['columns']])
                self.assertNotIn('excel_preview_error', response.data)
                self.assertEqual([sheet['name'] for sheet in response.data['excel_workbook_preview']['sheets']], register.parent.sheetnames)
            else:
                self.assertTrue(response.content.decode('utf-8-sig').lstrip('\ufeff').startswith('Timestamp,'))

    def test_mutation_has_evidence_even_when_reason_is_missing(self):
        log = self.log('rubric.update', old_values={'name': 'Old'}, new_values={'name': 'New'})
        self.assertEqual(log.review_status, 'requires_reason')
        wb = self.export()
        _, rows = self.data(wb['Audit Register'])
        self.assertEqual(rows[0]['Evidence Status'], 'Evidence captured')
        self.assertEqual(rows[0]['Reason for Change'], 'Reason required - not recorded')
        _, evidence = self.data(wb['Event Evidence'])
        self.assertEqual(evidence[0]['Before'], 'Old')
        self.assertEqual(evidence[0]['After'], 'New')

    def test_historical_review_never_invents_identity_or_timestamp(self):
        self.log(review_status='reviewed')
        _, rows = self.data(self.export()['Audit Register'])
        self.assertEqual(rows[0]['Review Status'], 'Reviewed')
        self.assertEqual(rows[0]['Reviewed By'], 'Reviewer not recorded')
        self.assertEqual(rows[0]['Reviewer Role'], 'Not recorded')
        self.assertEqual(rows[0]['Last Review Timestamp'], 'Not recorded')

    def test_reason_alone_and_scope_tags_do_not_claim_snapshot_evidence(self):
        self.log(new_values={'scope': 'capstone'}, reason='Configured a rubric')
        _, rows = self.data(self.export()['Audit Register'])
        self.assertEqual(rows[0]['Evidence Status'], 'Event details incomplete')

    def test_long_history_and_snapshot_values_are_preserved_as_literal_text(self):
        snapshot = '=HYPERLINK("https://example.invalid") ' + 'criteria detail ' * 3000
        note = 'Reviewed detailed evidence. ' * 1500
        log = self.log(new_values={'description': snapshot, 'enabled': False, 'score': 0, 'name': None})
        SystemAuditReview.objects.create(
            audit_log=log, reviewer=self.admin, reviewer_name='Auditor', reviewer_role='admin',
            from_status='captured', to_status='reviewed', reason=note,
        )
        wb = self.export()
        _, rows = self.data(wb['Event Evidence'])
        description = [row for row in rows if row['Recorded Field'] == 'description']
        self.assertEqual(''.join(row['After'] or '' for row in description), snapshot)
        self.assertEqual(next(row['After'] for row in rows if row['Recorded Field'] == 'enabled'), 'false')
        self.assertEqual(next(row['After'] for row in rows if row['Recorded Field'] == 'score'), '0')
        self.assertEqual(next(row['After'] for row in rows if row['Recorded Field'] == 'name'), 'null')
        _, reviews = self.data(wb['Review History'])
        self.assertEqual(''.join(row['Review Note'] or '' for row in reviews), note)
        self.assertTrue(all(cell.data_type != 'f' for ws in wb for row in ws for cell in row))
        self.assertTrue(all(dimension.height <= 409 for ws in wb for dimension in ws.row_dimensions.values()))

    def test_empty_workbook_and_limited_export_state_their_scope(self):
        wb = self.export()
        self.assertEqual(wb.sheetnames, ['Audit Register', 'Event Evidence', 'Review History'])
        self.assertTrue(any(cell.value == 'No records match the selected filters.' for row in wb.active for cell in row))
        self.log()
        self.log()
        wb = self.export(limit=1)
        self.assertTrue(any('Latest 1 of 2 matching events' in str(cell.value) for row in wb['Event Evidence'] for cell in row))

    def test_api_exposes_capture_status_separately_from_review_and_reason(self):
        log = self.log('rubric.delete', old_values={'name': 'Removed'})
        response = self.client.get('/api/audit-logs/')
        row = next(row for row in response.data['audit_logs'] if row['id'] == log.pk)
        self.assertEqual(row['evidence']['status'], 'captured')
        self.assertTrue(row['evidence']['reason_required'])
        self.assertEqual(row['review_status'], 'requires_reason')

