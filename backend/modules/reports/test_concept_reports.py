import io
from pathlib import Path

import openpyxl
from django.contrib.auth import get_user_model
from django.utils import timezone
from rest_framework.test import APITestCase

from academic_period_management.models import SchoolYear, Semester
from defense.scheduler.models import PitEventGradingConfig
from defense.stages.models import DefenseStage
from grading.grades.models import TeamGrade
from repository.deliverables.models import ConceptRound, DeliverableSubmission
from repository.section_concepts.models import SectionConceptProject, SectionConceptRound, SectionSystemConcept
from student_teams.models import StudentTeam, TeamMembership, SectionAssignment
from user_management.academic_records.models import StudentAcademicRecord

User = get_user_model()


class ConceptProjectReportTests(APITestCase):
    def setUp(self):
        self.admin = User.objects.create_user(username='register-admin', role='admin')
        self.adviser = User.objects.create_user(username='register-adviser', role='faculty', is_adviser=True)
        self.lead = User.objects.create_user(username='register-lead', role='faculty', is_pit_lead=True, pit_lead_year='3rd Year')
        year = SchoolYear.objects.create(label='2094-2095')
        self.source = Semester.objects.create(school_year=year, label=Semester.FIRST, is_active=True)
        self.target = Semester.objects.create(school_year=year, label=Semester.SECOND,
            capstone_team_creation_enabled=True, capstone_program_phase=Semester.PHASE_CAPSTONE_1)
        self.event = PitEventGradingConfig.objects.create(semester=self.source, year_level='3rd Year',
            event_name='PIT Concept', concept_config={'enabled': True, 'required_count': 2})
        self.team = self.make_team('Alpha', ['2094-001', '2094-002'])
        self.other = self.make_team('Beta', ['2094-003'])
        self.round = ConceptRound.objects.create(team=self.team, pit_event_config=self.event, verdict='approved',
            concepts=[{'key': 'c1', 'title': 'Campus Records', 'decision': 'approved'},
                      {'key': 'c2', 'title': 'Other Concept', 'decision': 'not_selected'}])
        self.params = {'semester_id': self.source.pk, 'scope': 'pit', 'year_level': '3rd Year',
            'pit_event_config_id': self.event.pk, 'export_format': 'json'}
        self.client.force_authenticate(self.admin)

    def make_team(self, name, students, level=StudentTeam.LEVEL_3_PIT, semester=None):
        students = [User.objects.create_user(username=s, role='student', first_name='Duplicate', last_name='Name') for s in students]
        team = StudentTeam.objects.create(name=name, project_title='Display title is not a decision',
            level=level, year_level=level.split(' PIT')[0].split(' Capstone')[0], semester=semester or self.source,
            section='BSIT-3A', leader=students[0], adviser=self.adviser)
        for index, student in enumerate(students):
            TeamMembership.objects.create(team=team, student=student, is_leader=index == 0)
            StudentAcademicRecord.objects.create(student=student, semester=team.semester, year_level=team.year_level, section=team.section)
        return team

    def get_report(self, **params):
        return self.client.get('/api/reports/concept-projects/', {**self.params, **params})

    def test_register_real_decisions_and_filters_do_not_create_domain_records(self):
        before = (ConceptRound.objects.count(), TeamGrade.all_objects.count(), DeliverableSubmission.all_objects.count())
        result = self.get_report()
        self.assertEqual(result.status_code, 200, result.data)
        self.assertEqual(len(result.data['rows']), 3)
        self.assertEqual([r['title'] for r in result.data['rows'] if r['concept_status'] == 'approved'], ['Campus Records'])
        self.assertEqual(len(self.get_report(concept_status='approved').data['rows']), 1)
        self.assertEqual(len(self.get_report(search='Other Concept').data['rows']), 1)
        self.assertEqual(before, (ConceptRound.objects.count(), TeamGrade.all_objects.count(), DeliverableSubmission.all_objects.count()))
        pdf = self.get_report(export_format='pdf')
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b'%PDF'))

    def test_new_round_and_project_version_hide_old_approval_by_default(self):
        ConceptRound.objects.create(team=self.team, pit_event_config=self.event, number=2,
            concepts=[{'key': 'c1', 'title': 'Replacement', 'decision': ''}])
        rows = self.get_report().data['rows']
        self.assertNotIn('Campus Records', [r['title'] for r in rows])
        self.assertFalse(any(r['eligible_title'] for r in rows))
        history = self.get_report(include_history='1').data['rows']
        old = next(r for r in history if r['title'] == 'Campus Records')
        self.assertFalse(old['current'])
        self.assertFalse(old['eligible_title'])
        self.team.project_version = 2
        self.team.save(update_fields=['project_version'])
        self.assertNotIn('Replacement', [r['title'] for r in self.get_report().data['rows']])

    def test_roster_uses_official_columns_stable_ids_and_leader_first(self):
        response = self.get_report(layout='roster', export_format='xlsx', include_signatures='true')
        self.assertEqual(response.status_code, 200, getattr(response, 'data', ''))
        sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
        self.assertEqual([c.value for c in sheet[3]], ['Team Name', 'Capstone Project', 'Section', 'Adviser', 'Team Members'])
        self.assertEqual(list(sheet.values)[3:6], [
            ('Alpha', 'Campus Records', 'BSIT-3A', None, '2094-001'),
            (None, None, None, None, '2094-002'),
            ('Beta', None, 'BSIT-3A', None, '2094-003')])
        self.assertIn('PIT Concept', sheet['A2'].value)
        # The Flutter parser regression reads these exact API-produced bytes.
        artifact = Path(__file__).resolve().parents[3] / '.tmp' / 'pit-title-roster.xlsx'
        artifact.parent.mkdir(exist_ok=True)
        artifact.write_bytes(response.content)

    def test_roster_requires_single_compatible_source_and_excel(self):
        for params in ({'semester_id': ''}, {'pit_event_config_id': ''}, {'scope': 'capstone'},
                {'concept_status': 'approved'}, {'include_history': '1'}, {'export_format': 'pdf'}):
            with self.subTest(params=params):
                self.assertEqual(self.get_report(layout='roster', **params).status_code, 400)

    def test_general_team_approval_does_not_export_a_title(self):
        self.other.status = StudentTeam.STATUS_APPROVED
        self.other.save(update_fields=['status'])
        TeamGrade.objects.create(team=self.other, semester=self.source, scope='pit', pit_event_config=self.event,
            final_grade=99, verdict='approved', stage_label=self.event.event_name, panel_weight=80, peer_weight=20, adviser_weight=0)
        rows = self.get_report(layout='roster').data['rows']
        self.assertEqual(next(r for r in rows if r['team'] == 'Beta')['title'], '')

    def test_revisions_need_matching_grade_clearance(self):
        self.round.verdict = 'approved_with_revisions'
        self.round.concepts = [{'key': 'c1', 'title': 'Campus Records', 'decision': 'approved_with_revisions'}]
        self.round.save()
        self.assertFalse(self.get_report().data['rows'][0]['eligible_title'])
        grade = TeamGrade.objects.create(team=self.team, semester=self.source, scope='pit', pit_event_config=self.event,
            verdict='approved_with_revisions', stage_label=self.event.event_name, panel_weight=80, peer_weight=20, adviser_weight=0)
        self.assertFalse(self.get_report().data['rows'][0]['eligible_title'])
        grade.revisions_cleared_at = timezone.now()
        grade.save(update_fields=['revisions_cleared_at'])
        self.assertTrue(self.get_report().data['rows'][0]['eligible_title'])

    def test_shared_module_requires_selected_system_and_its_clearance(self):
        section = SectionAssignment.objects.create(semester=self.source, year_level='3rd Year', section='BSIT-3A')
        project = SectionConceptProject.objects.create(section=section, pit_event_config=self.event,
            project_manager=self.team.leader, reviewer=self.adviser)
        shared = SectionConceptRound.objects.create(project=project, selected_key='c1', verdict='approved_with_revisions')
        SectionSystemConcept.objects.create(round=shared, key='c1', title='Shared System', decision='approved_with_revisions')
        self.round.section_round = shared
        self.round.save()
        self.assertFalse(self.get_report().data['rows'][0]['eligible_title'])
        shared.revisions_cleared_at = timezone.now()
        shared.save()
        self.assertTrue(self.get_report().data['rows'][0]['eligible_title'])
        shared.selected_key = 'c2'
        shared.save()
        self.assertFalse(self.get_report().data['rows'][0]['eligible_title'])

    def test_duplicate_membership_and_team_name_cannot_silently_merge(self):
        TeamMembership.objects.create(team=self.other, student=self.team.leader)
        self.assertEqual(self.get_report(layout='roster').status_code, 400)
        TeamMembership.objects.filter(team=self.other, student=self.team.leader).delete()
        self.other.name = self.team.name.lower()
        self.other.save(update_fields=['name'])
        self.assertEqual(self.get_report(layout='roster').status_code, 400)

    def test_lead_cannot_widen_reporting_via_adviser_role_or_filters(self):
        capstone = self.make_team('Capstone', ['2094-004'], StudentTeam.LEVEL_3_CAPSTONE, self.target)
        capstone.adviser = self.lead
        capstone.save()
        self.lead.is_adviser = True
        self.lead.save()
        self.client.force_authenticate(self.lead)
        self.assertEqual(self.get_report().status_code, 200)
        self.assertEqual(self.get_report(scope='capstone').status_code, 403)
        self.assertEqual(self.get_report(year_level='1st Year').status_code, 403)
        self.lead.pit_lead_year = ''
        self.lead.save()
        self.assertEqual(self.get_report().status_code, 403)
        self.client.force_authenticate(self.adviser)
        self.assertEqual(self.get_report().status_code, 403)

    def test_admin_reports_both_tracks_and_all_pit_years(self):
        for index, level in enumerate((StudentTeam.LEVEL_1_PIT, StudentTeam.LEVEL_2_PIT, StudentTeam.LEVEL_4_CAPSTONE)):
            self.make_team(f'Other {index}', [f'other-{index}'], level)
        result = self.get_report(scope='', year_level='', pit_event_config_id='')
        self.assertEqual({r['year_level'] for r in result.data['rows']}, {'1st Year', '2nd Year', '3rd Year', '4th Year'})
        self.assertEqual({r['scope'] for r in result.data['rows']}, {'PIT', 'Capstone'})

    def test_historical_source_remains_readable_for_third_year_lead(self):
        self.source.is_active = False
        self.source.save()
        self.target.is_active = True
        self.target.save()
        self.client.force_authenticate(self.lead)
        self.assertEqual(self.get_report(layout='roster', export_format='xlsx').status_code, 200)

    def test_roster_titles_are_literal_spreadsheet_text(self):
        self.round.concepts[0]['title'] = '=HYPERLINK("untrusted")'
        self.round.save()
        sheet = openpyxl.load_workbook(io.BytesIO(self.get_report(layout='roster', export_format='xlsx').content)).active
        self.assertEqual(sheet['B4'].data_type, 's')

    def test_preview_is_bounded_but_excel_contains_all_matching_rows(self):
        preview = self.get_report(page_size='1')
        self.assertEqual(len(preview.data['rows']), 1)
        self.assertEqual(preview.data['total_rows'], 3)
        self.assertTrue(preview.data['preview_limited'])
        self.assertIn('Preview Coverage', [m['label'] for m in preview.data['metadata']])
        sheet = openpyxl.load_workbook(io.BytesIO(self.get_report(export_format='xlsx', page_size='1').content)).active
        self.assertEqual(sheet.max_row, 6)

    def test_assigned_title_is_visible_before_submission_without_implying_approval(self):
        self.other.concept_assignment = {'mode': 'assigned', 'project_version': 1,
            'context_key': f'pit:{self.event.pk}', 'assigned_title': 'Assigned draft title'}
        self.other.save()
        row = next(r for r in self.get_report().data['rows'] if r['team'] == 'Beta')
        self.assertEqual(row['title'], 'Assigned draft title')
        self.assertEqual(row['concept_status'], 'not_submitted')
        self.assertFalse(row['eligible_title'])
        self.assertFalse(ConceptRound.objects.filter(team=self.other).exists())

    def test_export_then_admin_mixed_title_import_creates_fresh_capstone_records(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from django.test import override_settings
        from tempfile import TemporaryDirectory
        media = TemporaryDirectory()
        self.addCleanup(media.cleanup)
        storage = override_settings(MEDIA_ROOT=media.name)
        storage.enable()
        self.addCleanup(storage.disable)
        source_file = DeliverableSubmission.objects.create(team=self.team, stage_label=self.event.event_name,
            pit_event_config=self.event, deliverable_id='pit-only', label='PIT concept paper', deliverable_type='pre',
            status='accepted', reviewed_by=self.adviser, reviewed_at=timezone.now(), extracted_text='PIT concept',
            file=SimpleUploadedFile('pit-concept.pdf', b'%PDF-1.4 PIT source fixture'))
        source_grade = TeamGrade.objects.create(team=self.team, semester=self.source, scope='pit', pit_event_config=self.event,
            verdict='approved', stage_label=self.event.event_name, panel_weight=80, peer_weight=20, adviser_weight=0,
            revisions_cleared_at=timezone.now())
        ConceptRound.objects.create(team=self.other, pit_event_config=self.event, verdict='approved',
            concepts=[{'key': 'c1', 'title': 'PIT-only idea', 'decision': 'approved'}])
        roster = self.get_report(layout='roster').data['rows']
        before = self.round.concepts
        self.source.is_active = False
        self.source.save()
        self.target.is_active = True
        self.target.save()
        stage = DefenseStage.objects.order_by('display_order', 'pk').first()
        stage.concept_config = {'enabled': True, 'required_count': 2}
        stage.save()
        for membership in TeamMembership.objects.filter(team__in=[self.team, self.other]):
            StudentAcademicRecord.objects.create(student=membership.student, semester=self.target, year_level='3rd Year', section='BSIT-3A')
        rows = []
        for row in roster:
            if row['team']:
                rows.append({'team_name': row['team'], 'project_title': row['title'], 'section': row['section'],
                    'adviser_id': self.adviser.username, 'level': StudentTeam.LEVEL_3_CAPSTONE,
                    'member_ids': [], 'leader_id': row['members']})
            rows[-1]['member_ids'].append(row['members'])
        # Admin decides that Beta's approved PIT concept is unsuitable for Capstone.
        rows[1]['project_title'] = ''
        preview = self.client.post('/api/teams/bulk-import/preview/', {'teams': rows}, format='json')
        self.assertTrue(all(r['ready'] for r in preview.data['rows']), preview.data)
        self.assertEqual([r['concept_requirement']['required_count'] for r in preview.data['rows']], [1, 2])
        imported = self.client.post('/api/teams/bulk-import/', {'teams': rows}, format='json')
        self.assertEqual(imported.status_code, 201, imported.data)
        new = StudentTeam.objects.filter(semester=self.target)
        self.assertEqual(new.count(), 2)
        self.assertEqual({t.concept_assignment['mode'] for t in new}, {'assigned', 'proposals'})
        self.assertFalse(ConceptRound.objects.filter(team__in=new).exists())
        self.assertFalse(DeliverableSubmission.all_objects.filter(team__in=new).exists())
        self.assertFalse(TeamGrade.all_objects.filter(team__in=new, verdict__in=['approved', 'approved_with_revisions']).exists())
        self.assertTrue(all(t.status == StudentTeam.STATUS_PENDING for t in new))
        source_file.refresh_from_db()
        source_grade.refresh_from_db()
        self.assertEqual(source_file.status, 'accepted')
        self.assertEqual(source_file.team_id, self.team.pk)
        self.assertTrue(source_file.file.storage.exists(source_file.file.name))
        self.assertEqual(source_grade.verdict, 'approved')
        self.assertFalse(TeamGrade.all_objects.filter(team__in=new, revisions_cleared_at__isnull=False).exists())
        self.round.refresh_from_db()
        self.assertEqual(self.round.concepts, before)
        again = self.client.post('/api/teams/bulk-import/', {'teams': rows}, format='json')
        self.assertEqual(again.data['created_count'], 0)
