from django.contrib.auth import get_user_model
from rest_framework.test import APITestCase
from academic_period_management.models import SchoolYear, Semester
from student_teams.models import StudentTeam, TeamMembership
from grading.grades.models import TeamGrade
from defense.stages.models import DefenseStage

User = get_user_model()

class ReportsApiTests(APITestCase):
    def setUp(self):
        # 1. Users
        self.admin = User.objects.create_user(
            username='admin-user',
            password='pass12345',
            role='admin',
            is_staff=True,
        )
        self.adviser = User.objects.create_user(
            username='adviser-1',
            password='pass12345',
            role='faculty',
            first_name='Ada',
            last_name='Lovelace',
            is_adviser=True,
        )
        self.other_adviser = User.objects.create_user(
            username='adviser-2',
            password='pass12345',
            role='faculty',
            first_name='Alan',
            last_name='Turing',
            is_adviser=True,
        )
        self.pit_lead = User.objects.create_user(
            username='pit-lead-1',
            password='pass12345',
            role='faculty',
            is_pit_lead=True,
            pit_lead_year='3rd Year',
        )
        self.student = User.objects.create_user(
            username='2024-0001',
            password='pass12345',
            role='student',
        )

        # 2. Semester Setup
        self.school_year, _ = SchoolYear.objects.get_or_create(label='2026-2027')
        self.semester, _ = Semester.objects.get_or_create(
            school_year=self.school_year,
            label=Semester.FIRST,
            defaults={'is_active': True},
        )

        # 3. Defense Stage (needed for TeamGrade)
        self.stage, _ = DefenseStage.objects.get_or_create(
            label='Concept Proposal',
            defaults={'is_active': True, 'display_order': 1}
        )

        # 4. Student Teams
        self.team_advised = StudentTeam.objects.create(
            name='Team Advised',
            project_title='Adviser\'s Project',
            level=StudentTeam.LEVEL_3_PIT,
            year_level='3rd Year',
            semester=self.semester,
            leader=self.student,
            adviser=self.adviser,
        )
        self.team_other = StudentTeam.objects.create(
            name='Team Other',
            project_title='Other Project',
            level=StudentTeam.LEVEL_4_CAPSTONE,
            year_level='4th Year',
            semester=self.semester,
            leader=self.student,
            adviser=self.other_adviser,
        )

        # 5. Team Grades
        self.grade_advised = TeamGrade.objects.create(
            team=self.team_advised,
            semester=self.semester,
            scope=TeamGrade.SCOPE_PIT,
            stage_label='Concept Proposal',
            defense_stage=self.stage,
            panel_score=85,
            adviser_score=0,
            peer_score=88,
            final_grade=87,
            panel_weight=80,
            adviser_weight=0,
            peer_weight=20,
            status=TeamGrade.STATUS_PUBLISHED,
        )

        self.grade_other = TeamGrade.objects.create(
            team=self.team_other,
            semester=self.semester,
            scope=TeamGrade.SCOPE_CAPSTONE,
            stage_label='Concept Proposal',
            defense_stage=self.stage,
            panel_score=75,
            adviser_score=80,
            peer_score=78,
            final_grade=77,
            panel_weight=50,
            adviser_weight=30,
            peer_weight=20,
            status=TeamGrade.STATUS_PUBLISHED,
        )

    def test_formatted_grade_workbooks_remove_repeated_context_and_duplicate_totals(self):
        import base64
        import io
        from pathlib import Path
        import openpyxl
        import pdfplumber
        self.client.force_authenticate(user=self.admin)
        output = Path(__file__).resolve().parents[3] / '.tmp'
        output.mkdir(exist_ok=True)
        self.student.first_name, self.student.last_name = 'Adrian', 'Valdez'
        self.student.save(update_fields=['first_name', 'last_name'])
        # This test checks numeric workbook layout, so supply an actual individual
        # assessment rather than inheriting a team final for a missing record.
        from grading.grades.models import StudentStageGrade
        StudentStageGrade.objects.update_or_create(team_grade=self.grade_other, student=self.student,
            defaults={'panel_score': self.grade_other.panel_score, 'adviser_score': self.grade_other.adviser_score,
                      'peer_score': self.grade_other.peer_score, 'final_grade': self.grade_other.final_grade})
        for index, (first, last) in enumerate((('Stephanie', 'Yap'), ('Jerome', 'De Leon'), ('Camille', 'Roxas')), 1):
            member = User.objects.create_user(username=f'workbook-preview-student-{index}', role='student', first_name=first, last_name=last)
            TeamMembership.objects.create(team=self.team_other, student=member)
        for kind, record_id in (('team-grade', self.team_other.pk), ('individual-grade', self.student.pk)):
            url = f'/api/reports/{kind}/{record_id}/'
            params = {'scope': 'capstone', 'semester_id': self.semester.pk}
            preview = self.client.get(url, {**params, 'export_format': 'json'})
            export = self.client.get(url, {**params, 'export_format': 'xlsx'})
            self.assertEqual(export.status_code, 200)
            self.assertEqual(preview.status_code, 200)
            self.assertNotIn('excel_preview_error', preview.data)
            workbook = openpyxl.load_workbook(io.BytesIO(export.content))
            self.assertEqual([sheet['name'] for sheet in preview.data['excel_workbook_preview']['sheets']], workbook.sheetnames)
            sheet = workbook.active
            header = openpyxl.utils.range_boundaries(sheet.auto_filter.ref)[1]
            labels = [cell.value for cell in sheet[header]]
            for repeated in ('Scope', 'Stage / Event', 'Academic Period', 'Team', 'Panel Weight', 'Panel Evaluators'):
                self.assertNotIn(repeated, labels)
            self.assertGreater(header, 3)
            self.assertFalse(sheet.protection.sheet)
            self.assertEqual(sheet.freeze_panes, f'A{header + 1}')
            if kind == 'team-grade':
                values = dict(zip(labels, [cell.value for cell in sheet[header + 1]]))
                self.assertEqual(values['Panel Contribution\n(50%)'], 37.5)
                self.assertEqual(values['Total Score'], float(self.grade_other.final_grade))
                self.assertEqual(values['Grade Status'], 'Grade passed')
                self.assertEqual(sum(label in labels for label in ('Total Score', 'Final Grade')), 1)
            else:
                self.assertIn('Weighted Contribution', labels)
                data_rows = list(sheet.iter_rows(min_row=header + 1))
                self.assertEqual(data_rows[0][3].value, 37.5)
                self.assertEqual(data_rows[-1][0].value, 'TOTAL OFFICIAL GRADE')
            pdf_bytes = base64.b64decode(preview.data['excel_workbook_preview']['sheets'][0]['pdf_base64'])
            with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
                text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
                self.assertTrue(all(char['top'] >= 0 for page in pdf.pages for char in page.chars))
            self.assertIn('Other Project', text)
            self.assertIn(f'{self.grade_other.final_grade:.2f}', text)
            (output / f'{kind}-workbook-qa.xlsx').write_bytes(export.content)
            (output / f'{kind}-workbook-preview-qa.pdf').write_bytes(pdf_bytes)

    def test_all_stage_workbook_has_separate_assessment_sheets(self):
        import io
        import openpyxl
        stage, _ = DefenseStage.objects.get_or_create(label='Project Proposal', defaults={'display_order': 2})
        TeamGrade.objects.create(team=self.team_other, semester=self.semester, scope='capstone', stage_label=stage.label,
                                 defense_stage=stage, panel_score=80, adviser_score=85, peer_score=90, final_grade=84,
                                 panel_weight=50, adviser_weight=30, peer_weight=20)
        self.client.force_authenticate(user=self.admin)
        response = self.client.get(f'/api/reports/team-grade/{self.team_other.pk}/', {'coverage': 'all', 'export_format': 'xlsx'})
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ['Concept Proposal', 'Project Proposal'])
        for sheet in workbook:
            metadata = [cell.value for row in sheet for cell in row]
            self.assertIn(sheet.title, metadata)
            self.assertIn('All recorded stages / events', metadata)

    def test_workbook_preserves_criteria_long_feedback_and_attempt_history(self):
        import io
        import openpyxl
        from datetime import date, time
        from defense.scheduler.models import DefenseSchedule
        from grading.rubrics.models import Rubric, RubricCriterion
        from grading.grades.models import PanelistGradeSubmission, PanelistCriterionScore, GradeAttemptHistory
        rubric = Rubric.objects.create(name='Test Scoring Guide', scope='capstone', semester=self.semester,
                                       defense_stage=self.stage, evaluation_type='panel', created_by=self.admin)
        criterion = RubricCriterion.objects.create(rubric=rubric, name='System Design', weight=100, max_score=10, scale=Rubric.SCALE_10)
        schedule = DefenseSchedule.objects.create(team=self.team_other, semester=self.semester, defense_stage=self.stage,
                                                  scope='capstone', scheduled_date=date(2026, 6, 18), start_time=time(8), room='Room A')
        remarks = '=Literal feedback. ' * 100 + 'COMPLETE FEEDBACK END'
        submission = PanelistGradeSubmission.objects.create(team_grade=self.grade_other, schedule=schedule,
                                                             panelist=self.other_adviser, remarks=remarks)
        PanelistCriterionScore.objects.create(submission=submission, criterion=criterion, criterion_name_snapshot=criterion.name,
                                              score=9, max_score_snapshot=10)
        GradeAttemptHistory.objects.create(team_grade=self.grade_other, attempt_number=1, schedule=schedule,
                                           panel_score=70, final_grade=70, verdict='approved', verdict_remarks='First attempt')
        self.client.force_authenticate(user=self.admin)
        url = f'/api/reports/team-grade/{self.team_other.pk}/'
        export = self.client.get(url, {'export_format': 'xlsx'})
        self.assertEqual(export.status_code, 200)
        workbook = openpyxl.load_workbook(io.BytesIO(export.content))
        feedback = next(sheet for sheet in workbook if sheet.title.endswith('Feedback'))
        header = openpyxl.utils.range_boundaries(feedback.auto_filter.ref)[1]
        self.assertEqual(''.join(row[3].value for row in feedback.iter_rows(min_row=header + 1)), remarks)
        self.assertFalse(any(cell.data_type == 'f' for sheet in workbook for row in sheet for cell in row))
        history = next(sheet for sheet in workbook if sheet.title.endswith('History'))
        header = openpyxl.utils.range_boundaries(history.auto_filter.ref)[1]
        self.assertEqual(history.cell(header + 1, 2).value.date(), date(2026, 6, 18))
        self.assertEqual(history.cell(header + 1, 5).value, 'Approved')
        preview = self.client.get(url, {'export_format': 'json'})
        self.assertEqual(preview.status_code, 200)
        self.assertTrue(preview.data['pdf_base64'])
        self.assertNotIn('excel_preview_error', preview.data)
        self.assertEqual([sheet['name'] for sheet in preview.data['excel_workbook_preview']['sheets']], workbook.sheetnames)

    def test_anonymous_access_denied(self):
        endpoints = [
            f'/api/reports/team-grade/{self.team_advised.id}/',
            '/api/reports/semester-grades/',
            '/api/reports/defense-schedules/',
            '/api/reports/team-roster/',
            '/api/reports/user-directory/',
            '/api/reports/audit-trail/',
        ]
        for url in endpoints:
            response = self.client.get(url)
            self.assertEqual(response.status_code, 401)

    def test_admin_can_generate_all_reports(self):
        self.client.force_authenticate(user=self.admin)
        
        # Team Grade Report
        url = f'/api/reports/team-grade/{self.team_advised.id}/'
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertIn('attachment', response['Content-Disposition'])

        # Semester Grades
        response = self.client.get('/api/reports/semester-grades/', {'semester_id': self.semester.id})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')

        # Defense Schedules
        response = self.client.get('/api/reports/defense-schedules/', {'semester_id': self.semester.id})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')

        # Team Roster
        response = self.client.get('/api/reports/team-roster/', {'semester_id': self.semester.id})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')

        # User Directory
        response = self.client.get('/api/reports/user-directory/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')

        # Audit Trail
        response = self.client.get('/api/reports/audit-trail/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')

    def test_team_grade_signature_settings_match_preview_and_export(self):
        import base64
        import io
        import json
        import pdfplumber

        self.client.force_authenticate(user=self.admin)
        url = f'/api/reports/team-grade/{self.team_other.id}/'
        signers = [{'label': 'Verified by:', 'name': 'Report Reviewer', 'role': 'Quality Officer'}]
        for include in ('true', 'false', None):
            params = {'signatories': json.dumps(signers)}
            if include is not None:
                params['include_signatures'] = include
            preview = self.client.get(url, {**params, 'export_format': 'json'})
            export = self.client.get(url, params)
            self.assertEqual(preview.status_code, 200)
            self.assertEqual(export.status_code, 200)
            self.assertTrue(preview.data.get('pdf_base64'))

            for source, pdf_bytes in (
                ('preview', base64.b64decode(preview.data['pdf_base64'])),
                ('export', export.content),
            ):
                with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
                    text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
                with self.subTest(include_signatures=include, source=source):
                    self.assertIn('Master Student Grade Sheet', text)
                    for signature_text in ('Verified by:', 'Report Reviewer', 'Quality Officer'):
                        if include == 'true':
                            self.assertIn(signature_text, text)
                        else:
                            self.assertNotIn(signature_text, text)
                    self.assertNotIn('Prepared by:', text)
                    self.assertNotIn('Approved by:', text)

    def test_all_excel_reports_use_the_selected_signatories(self):
        import io
        import json
        import openpyxl

        self.client.force_authenticate(user=self.admin)
        signers = [{'label': 'Verified by:', 'name': 'Custom Report Reviewer', 'role': 'Quality Officer'}]
        endpoints = [
            f'/api/reports/team-grade/{self.team_other.pk}/',
            f'/api/reports/individual-grade/{self.student.pk}/',
            '/api/reports/semester-grades/',
            '/api/reports/defense-schedules/',
            '/api/reports/team-roster/',
            '/api/reports/user-directory/',
        ]
        for endpoint in endpoints:
            for enabled in ('true', 'false'):
                with self.subTest(endpoint=endpoint, include_signatures=enabled):
                    response = self.client.get(endpoint, {
                        'export_format': 'xlsx', 'semester_id': self.semester.pk,
                        'include_signatures': enabled, 'signatories': json.dumps(signers),
                    })
                    self.assertEqual(response.status_code, 200)
                    self.assertIn('spreadsheetml', response['Content-Type'])
                    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
                    sheet = workbook.active
                    values = [cell.value for row in sheet for cell in row]
                    for text in ('Verified by:', 'Custom Report Reviewer', 'Quality Officer'):
                        self.assertEqual(text in values, enabled == 'true')
                    self.assertNotIn('Prepared by:', values)
                    if enabled == 'true':
                        data_end = openpyxl.utils.range_boundaries(sheet.auto_filter.ref)[3]
                        signer = next(cell for row in sheet for cell in row if cell.value == 'Custom Report Reviewer')
                        self.assertGreater(signer.row, data_end)
                        self.assertIn(f'${sheet.max_row}', sheet.print_area)

    def test_semester_report_explains_approval_below_grade_threshold(self):
        import base64
        import io
        from decimal import Decimal
        import pdfplumber

        self.client.force_authenticate(user=self.admin)
        self.grade_other.panel_score = Decimal('58.76')
        self.grade_other.adviser_score = Decimal('85.00')
        self.grade_other.peer_score = Decimal('73.30')
        self.grade_other.verdict = TeamGrade.VERDICT_APPROVED
        self.grade_other.save(update_fields=[
            'panel_score', 'adviser_score', 'peer_score', 'final_grade', 'verdict',
        ])

        params = {'semester_id': self.semester.id, 'scope': 'capstone'}
        preview = self.client.get('/api/reports/semester-grades/', {**params, 'export_format': 'json'})
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.data['rows'][0]['final_grade'], '69.54%')
        self.assertEqual(preview.data['rows'][0]['result'], 'Approved')
        metadata = {item['label']: item['value'] for item in preview.data['metadata']}
        self.assertIn('threshold is 75%', metadata['Outcome Basis'])
        self.assertIn('average is for reference', metadata['Outcome Basis'])

        export = self.client.get('/api/reports/semester-grades/', params)
        self.assertEqual(export.status_code, 200)
        for pdf_bytes in (export.content, base64.b64decode(preview.data['pdf_base64'])):
            with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
                text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
            self.assertIn('69.54%', text)
            self.assertIn('Below 75%', text)
            self.assertIn('Approved', text)
            self.assertIn('Concept Proposal - Capstone - 1 record', text)

        # A report must keep the no-verdict grade fallback distinguishable too.
        self.grade_other.verdict = ''
        self.grade_other.save(update_fields=['verdict'])
        preview = self.client.get('/api/reports/semester-grades/', {**params, 'export_format': 'json'})
        self.assertEqual(preview.data['rows'][0]['result'], 'Grade failed')

    def test_semester_report_counts_unique_teams_and_all_stage_outcomes(self):
        from decimal import Decimal

        self.client.force_authenticate(user=self.admin)
        self.grade_other.verdict = TeamGrade.VERDICT_APPROVED
        self.grade_other.save(update_fields=['verdict'])
        for index, (label, verdict, final_grade) in enumerate([
            ('Project Proposal', TeamGrade.VERDICT_APPROVED_WITH_REVISIONS, Decimal('77.07')),
            ('Pre-Oral Defense', TeamGrade.VERDICT_FOR_REDEFENSE, Decimal('70.00')),
            ('Final Defense', TeamGrade.VERDICT_PROJECT_REJECTED, Decimal('68.00')),
            ('Additional Evaluation', '', None),
        ], start=2):
            stage, _ = DefenseStage.objects.get_or_create(
                label=label, defaults={'display_order': index},
            )
            TeamGrade.objects.create(
                team=self.team_other, semester=self.semester,
                scope=TeamGrade.SCOPE_CAPSTONE, defense_stage=stage,
                stage_label=label, final_grade=final_grade, verdict=verdict,
                panel_score=final_grade, adviser_score=final_grade, peer_score=final_grade,
            )

        preview = self.client.get('/api/reports/semester-grades/', {
            'semester_id': self.semester.id, 'export_format': 'json',
        })
        self.assertEqual(preview.status_code, 200)
        metadata = {item['label']: item['value'] for item in preview.data['metadata']}
        self.assertEqual(metadata['Teams in Report'], '2')
        self.assertEqual(metadata['Stage / Event Records'], '6')
        self.assertEqual(metadata['Outcome Breakdown'], (
            'Passed / approved: 2; Failed / rejected: 1; Revisions pending: 1; '
            'For re-defense: 1; Pending grades: 1'
        ))
        outcomes = {row['result'] for row in preview.data['rows']}
        self.assertIn('Revisions pending', outcomes)
        self.assertIn('For Re-defense', outcomes)
        self.assertIn('Project Rejected', outcomes)
        self.assertIn('Pending grades', outcomes)

    def test_semester_pdf_repeats_stage_and_column_headers_on_continuation_pages(self):
        import io
        from decimal import Decimal
        import pdfplumber
        from reports.generators.semester_grades_report import generate_semester_grades_pdf

        # Unsaved records exercise pagination without writing extra team data.
        records = [TeamGrade(
            team=self.team_other, semester=self.semester,
            scope=TeamGrade.SCOPE_CAPSTONE, defense_stage=self.stage,
            stage_label=self.stage.label, final_grade=Decimal('77.10'),
            panel_score=Decimal('75.00'), adviser_score=Decimal('80.00'),
            peer_score=Decimal('78.00'), verdict=TeamGrade.VERDICT_APPROVED,
        ) for _ in range(60)]
        pdf_bytes = generate_semester_grades_pdf(self.semester, records, 'Report Test')
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            self.assertGreater(len(pdf.pages), 1)
            for page in pdf.pages:
                text = page.extract_text() or ''
                self.assertIn('Concept Proposal - Capstone - 60 records', text)
                self.assertIn('Team Name', text)
                self.assertIn('Adviser', text)
                self.assertIn('Outcome', text)

    def test_team_report_supports_latest_specific_and_all_stage_coverage(self):
        import base64
        import io
        import pdfplumber
        from decimal import Decimal

        self.client.force_authenticate(user=self.admin)
        stage, _ = DefenseStage.objects.get_or_create(label='Project Proposal', defaults={'display_order': 2})
        later = TeamGrade.objects.create(
            team=self.team_other, semester=self.semester, scope='capstone',
            defense_stage=stage, panel_score=Decimal('80'), adviser_score=Decimal('80'),
            peer_score=Decimal('80'),
        )
        url = f'/api/reports/team-grade/{self.team_other.pk}/'
        params = {'scope': 'capstone', 'semester_id': self.semester.pk, 'export_format': 'json'}
        latest = self.client.get(url, params)
        self.assertEqual(latest.status_code, 200)
        self.assertEqual({row['assessment'] for row in latest.data['rows']}, {'Project Proposal'})
        self.assertEqual(len(latest.data['assessments']), 2)
        selected = self.client.get(url, {**params, 'coverage': 'specific', 'grade_id': self.grade_other.pk})
        self.assertEqual({row['assessment'] for row in selected.data['rows']}, {'Concept Proposal'})
        all_stages = self.client.get(url, {**params, 'coverage': 'all'})
        self.assertEqual({row['assessment'] for row in all_stages.data['rows']}, {'Concept Proposal', 'Project Proposal'})
        with pdfplumber.open(io.BytesIO(base64.b64decode(all_stages.data['pdf_base64']))) as pdf:
            text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
            self.assertIn('Concept Proposal', text)
            self.assertIn('Project Proposal', text)
            self.assertNotIn('PROJECT FINAL DEFENSE', text)
            self.assertGreaterEqual(len(pdf.pages), 2)
        mismatch = self.client.get(url, {**params, 'grade_id': self.grade_advised.pk})
        self.assertEqual(mismatch.status_code, 404)
        mismatch = self.client.get(url, {**params, 'scope': 'pit', 'grade_id': later.pk})
        self.assertEqual(mismatch.status_code, 404)

    def test_individual_report_selects_scope_from_assessment_not_first_team(self):
        import base64
        import io
        import pdfplumber

        self.client.force_authenticate(user=self.admin)
        url = f'/api/reports/individual-grade/{self.student.pk}/'
        for scope, team, label in [('pit', self.team_advised, 'PIT'), ('capstone', self.team_other, 'Capstone')]:
            response = self.client.get(url, {
                'scope': scope, 'semester_id': self.semester.pk, 'coverage': 'all', 'export_format': 'json',
            })
            self.assertEqual(response.status_code, 200)
            self.assertEqual({row['scope'] for row in response.data['rows']}, {label})
            self.assertEqual({row['team_name'] for row in response.data['rows']}, {team.name})
            with pdfplumber.open(io.BytesIO(base64.b64decode(response.data['pdf_base64']))) as pdf:
                text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
            self.assertIn(label.upper() + ' INDIVIDUAL GRADE AUDIT REPORT', text)
            if scope == 'pit':
                self.assertIn('PIT Event', text)
                self.assertNotIn('Capstone Defense Stage', text)
        csv = self.client.get(url, {'coverage': 'all', 'export_format': 'csv'})
        self.assertIn('PIT', csv.content.decode('utf-8-sig'))
        self.assertIn('Capstone', csv.content.decode('utf-8-sig'))

    def test_report_candidates_preserve_student_context_and_visibility(self):
        self.client.force_authenticate(user=self.admin)
        url = '/api/reports/assessment-candidates/'
        response = self.client.get(url, {'scope': 'pit', 'semester_id': self.semester.pk})
        self.assertEqual(response.status_code, 200)
        self.assertEqual([team['id'] for team in response.data['teams']], [self.team_advised.pk])
        context = response.data['student_contexts'][str(self.student.pk)]['contexts'][0]
        self.assertEqual(context['team']['scope'], 'pit')
        self.client.force_authenticate(user=self.adviser)
        response = self.client.get(url)
        self.assertEqual({team['id'] for team in response.data['teams']}, {self.team_advised.pk})

    def test_individual_and_candidate_selection_respects_academic_period(self):
        from decimal import Decimal
        self.client.force_authenticate(user=self.admin)
        period, _ = Semester.objects.get_or_create(school_year=self.school_year, label=Semester.SECOND)
        historical_team = StudentTeam.objects.create(
            name='PIT Other Period', project_title='Historical PIT Assessment',
            level=StudentTeam.LEVEL_2_PIT, year_level='2nd Year',
            semester=period, leader=self.student, adviser=self.adviser,
        )
        TeamGrade.objects.create(
            team=historical_team, semester=period, scope='pit', stage_label='Year 2 Showcase',
            panel_score=Decimal('80'), peer_score=Decimal('85'),
            panel_weight=80, adviser_weight=0, peer_weight=20,
        )
        params = {'scope': 'pit', 'semester_id': period.pk}
        candidates = self.client.get('/api/reports/assessment-candidates/', params)
        self.assertEqual([team['id'] for team in candidates.data['teams']], [historical_team.pk])
        report = self.client.get(f'/api/reports/individual-grade/{self.student.pk}/', {
            **params, 'coverage': 'all', 'export_format': 'json',
        })
        self.assertEqual(report.status_code, 200)
        self.assertEqual({row['team_name'] for row in report.data['rows']}, {historical_team.name})
        self.assertEqual({row['assessment'] for row in report.data['rows']}, {'Year 2 Showcase'})

    def test_scope_is_explicit_in_mixed_summary_and_roster_data(self):
        self.client.force_authenticate(user=self.admin)
        for endpoint in ('semester-grades', 'team-roster'):
            response = self.client.get(f'/api/reports/{endpoint}/', {
                'semester_id': self.semester.pk, 'export_format': 'json',
            })
            self.assertEqual(response.status_code, 200)
            self.assertEqual({row['scope'] for row in response.data['rows']}, {'Capstone', 'PIT'})
            scoped = self.client.get(f'/api/reports/{endpoint}/', {
                'semester_id': self.semester.pk, 'scope': 'pit', 'export_format': 'json',
            })
            self.assertEqual({row['scope'] for row in scoped.data['rows']}, {'PIT'})

    def test_missing_required_panel_grade_is_pending_in_team_data_and_pdf(self):
        import base64
        import io
        import pdfplumber

        self.client.force_authenticate(user=self.admin)
        self.grade_other.status = TeamGrade.STATUS_PENDING
        self.grade_other.panel_score = None
        self.grade_other.save(update_fields=['status', 'panel_score', 'final_grade'])
        response = self.client.get(f'/api/reports/team-grade/{self.team_other.pk}/', {'export_format': 'json'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['rows'][0]['final_grade'], 'Pending')
        self.assertEqual(response.data['rows'][0]['grade_status'], 'Pending grading')
        with pdfplumber.open(io.BytesIO(base64.b64decode(response.data['pdf_base64']))) as pdf:
            text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
            cells = [' '.join((cell or '').split()) for page in pdf.pages for table in page.extract_tables() for row in table for cell in row]
        self.assertIn('Pending grading', cells)
        self.assertNotIn('FAILED', text)

    def test_signatures_off_omits_certification_and_extra_page(self):
        import io
        import pdfplumber
        from reports.pdf_builder import DefensysPdfReportBuilder

        page_counts = {}
        for include in (True, False):
            builder = DefensysPdfReportBuilder(title='Signature Toggle Regression')
            builder.add_header()
            # Leave too little room for the signature block to fit on page one.
            builder.add_spacer(6.4)
            builder.add_paragraph('End of report data')
            builder.add_signatures(
                include_signatures=include,
                certification_text='Certification regression marker',
            )
            with pdfplumber.open(io.BytesIO(builder.build())) as pdf:
                page_counts[include] = len(pdf.pages)
                text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
            if include:
                self.assertIn('Certification regression marker', text)
                self.assertIn('Prepared by:', text)
            else:
                self.assertNotIn('Certification regression marker', text)
                self.assertNotIn('Prepared by:', text)
        self.assertEqual(page_counts[False], 1)
        self.assertGreater(page_counts[True], page_counts[False])

    def test_adviser_access_limits(self):
        self.client.force_authenticate(user=self.adviser)

        # Team Grade Report - team they advise
        url = f'/api/reports/team-grade/{self.team_advised.id}/'
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')

        # Team Grade Report - team they do not advise
        url = f'/api/reports/team-grade/{self.team_other.id}/'
        response = self.client.get(url)
        self.assertEqual(response.status_code, 403)

        # User Directory (adviser is not admin)
        response = self.client.get('/api/reports/user-directory/')
        self.assertEqual(response.status_code, 403)

        # Audit Trail (adviser cannot review audit logs)
        response = self.client.get('/api/reports/audit-trail/')
        self.assertEqual(response.status_code, 403)

    def test_pit_lead_access(self):
        self.client.force_authenticate(user=self.pit_lead)

        # Audit Trail - PIT Leads have permission to audit trail (if can_review_audit_logs is True)
        response = self.client.get('/api/reports/audit-trail/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')

    def test_pit_lead_report_center_is_limited_to_assigned_year(self):
        import datetime
        from defense.scheduler.models import DefenseSchedule
        from authentication_access_control.models import SystemAuditLog

        teams = {'3rd Year': self.team_advised}
        for number, level, year in (
            (1, StudentTeam.LEVEL_1_PIT, '1st Year'),
            (2, StudentTeam.LEVEL_2_PIT, '2nd Year'),
        ):
            team = StudentTeam.objects.create(
                name=f'PIT Year {number} Team', project_title=f'Year {number} Project',
                level=level, year_level=year, semester=self.semester,
                leader=self.student, adviser=self.adviser,
            )
            TeamGrade.objects.create(
                team=team, semester=self.semester, scope='pit',
                stage_label=f'PIT Year {number} Event', panel_score=80,
                peer_score=85, panel_weight=80, adviser_weight=0, peer_weight=20,
            )
            teams[year] = team
        for year, team in teams.items():
            DefenseSchedule.objects.create(
                scope='pit', team=team, semester=self.semester,
                event_name=f'{year} PIT Event',
                scheduled_date=datetime.date(2026, 10, 15),
                start_time=datetime.time(9, 0), room='Room 1', created_by=self.admin,
            )
            SystemAuditLog.objects.create(
                actor=self.admin, category=SystemAuditLog.CATEGORY_REPOSITORY,
                action=f'repository.{year.replace(" ", "_")}',
                target_type='VaultEntry', target_id=str(team.pk),
                new_values={'entry_type': 'pit', 'year_level': year},
            )
        SystemAuditLog.objects.create(
            actor=self.admin, category=SystemAuditLog.CATEGORY_REPOSITORY,
            action='repository.unassigned_pit', target_type='VaultEntry', target_id='0',
            new_values={'entry_type': 'pit'},
        )
        SystemAuditLog.objects.create(
            actor=self.admin, category=SystemAuditLog.CATEGORY_REPOSITORY,
            action='repository.cross_year_pit', target_type='VaultEntry', target_id='0',
            old_values={'entry_type': 'pit', 'year_level': '1st Year'},
            new_values={'entry_type': 'pit', 'year_level': '2nd Year'},
        )

        for year, team in teams.items():
            with self.subTest(year=year):
                self.pit_lead.pit_lead_year = year
                self.pit_lead.save(update_fields=['pit_lead_year'])
                self.client.force_authenticate(user=self.pit_lead)
                candidates = self.client.get('/api/reports/assessment-candidates/')
                self.assertEqual({item['id'] for item in candidates.data['teams']}, {team.pk})
                for path, key in (
                    ('semester-grades', 'team_name'),
                    ('team-roster', 'team_name'),
                    ('defense-schedules', 'team_name'),
                ):
                    response = self.client.get(f'/api/reports/{path}/', {
                        'semester_id': self.semester.pk, 'export_format': 'json',
                    })
                    self.assertEqual(response.status_code, 200)
                    self.assertEqual({row[key] for row in response.data['rows']}, {team.name})
                    self.assertEqual({row['scope'] for row in response.data['rows']}, {'PIT'})
                individual = self.client.get(f'/api/reports/individual-grade/{self.student.pk}/', {
                    'coverage': 'all', 'export_format': 'json',
                })
                self.assertEqual({row['team_name'] for row in individual.data['rows']}, {team.name})
                own = self.client.get(f'/api/reports/team-grade/{team.pk}/', {'export_format': 'json'})
                self.assertEqual(own.status_code, 200)
                other_team = next(value for value in teams.values() if value.pk != team.pk)
                denied = self.client.get(f'/api/reports/team-grade/{other_team.pk}/', {'export_format': 'json'})
                self.assertEqual(denied.status_code, 403)
                log = self.client.get('/api/reports/audit-trail/', {'export_format': 'json'})
                self.assertEqual({row['action'] for row in log.data['rows']}, {
                    f'repository.{year.replace(" ", "_")}',
                })
                self.assertEqual(self.client.get('/api/reports/user-directory/').status_code, 403)

        self.pit_lead.pit_lead_year = ''
        self.pit_lead.save(update_fields=['pit_lead_year'])
        self.client.force_authenticate(user=self.pit_lead)
        roster = self.client.get('/api/reports/team-roster/', {'export_format': 'json'})
        self.assertEqual(roster.data['rows'], [])
        candidates = self.client.get('/api/reports/assessment-candidates/')
        self.assertEqual(candidates.data['teams'], [])

    def test_report_search_filters(self):
        self.client.force_authenticate(user=self.admin)
        
        # Create a panelist user
        panelist = User.objects.create_user(
            username='panelist-1',
            password='pass12345',
            role='faculty',
            first_name='Charles',
            last_name='Babbage',
            is_panelist=True,
        )
        
        # Create a defense schedule
        from defense.scheduler.models import DefenseSchedule, SchedulePanelist
        import datetime
        schedule = DefenseSchedule.objects.create(
            scope=DefenseSchedule.SCOPE_CAPSTONE,
            semester=self.semester,
            team=self.team_other,
            defense_stage=self.stage,
            scheduled_date=datetime.date(2026, 7, 4),
            start_time=datetime.time(10, 0),
            room='Room 101',
            created_by=self.admin,
        )
        
        # Assign panelist to schedule
        SchedulePanelist.objects.create(
            schedule=schedule,
            panelist=panelist,
            order=1,
            is_chair=True,
        )
        
        # Link schedule to TeamGrade
        self.grade_other.schedule = schedule
        self.grade_other.save()
        
        # 1. Test Semester Grades search
        # Search by team adviser first name 'Ada' (should match team_advised)
        response = self.client.get('/api/reports/semester-grades/', {
            'semester_id': self.semester.id,
            'search': 'Ada'
        })
        self.assertEqual(response.status_code, 200)
        
        # Search by panelist last name 'Babbage'
        response = self.client.get('/api/reports/semester-grades/', {
            'semester_id': self.semester.id,
            'search': 'Babbage'
        })
        self.assertEqual(response.status_code, 200)

        # 2. Test Defense Schedules search
        # Search by stage label 'Concept'
        response = self.client.get('/api/reports/defense-schedules/', {
            'semester_id': self.semester.id,
            'search': 'Concept'
        })
        self.assertEqual(response.status_code, 200)

        # Search by panelist first name 'Charles'
        response = self.client.get('/api/reports/defense-schedules/', {
            'semester_id': self.semester.id,
            'search': 'Charles'
        })
        self.assertEqual(response.status_code, 200)

    def test_defensys_pdf_report_builder_direct(self):
        from reports.pdf_builder import DefensysPdfReportBuilder
        from reportlab.lib.units import inch

        builder = DefensysPdfReportBuilder(
            title="System Unit Test Report",
            subtitle="Verification of Centralized PDF Builder",
            generated_by="Automated Tester",
        )
        builder.add_header()
        builder.add_metadata_grid([
            ("Test Case Name", "DefensysPdfReportBuilder Direct Generation"),
            ("Status", "Running Test"),
            ("Date", "2026-09-01"),
        ])
        builder.add_section_header("Section 1: Data Table Verification")
        builder.add_paragraph("This paragraph verifies typography and text wrapping in the centralized PDF builder.")
        builder.add_table(
            headers=["ID", "Description", "Result", "Score"],
            rows=[
                ["TC-001", "Header rendering verification", "PASSED", "100%"],
                ["TC-002", "Numbered canvas calculation", "PASSED", "100%"],
                ["TC-003", "Zebra-striped table rows", "PASSED", "100%"],
            ],
            col_widths=[1.0 * inch, 3.5 * inch, 1.2 * inch, 1.0 * inch],
            alignments=['center', 'left', 'center', 'center'],
            bold_cols=[0, 2],
        )
        builder.add_alert_box(
            title="Quality Verification Notice",
            message="All styling rules conform to official USTP Department of Information Technology guidelines.",
            level="success",
        )
        builder.add_signatures(
            prepared_by="Automated Test Runner",
            noted_by="Quality Assurance Officer",
            approved_by="IT Program Chairperson",
        )

        pdf_bytes = builder.build()
        self.assertIsInstance(pdf_bytes, bytes)
        self.assertTrue(len(pdf_bytes) > 1000)
        self.assertTrue(pdf_bytes.startswith(b'%PDF'))

    def test_defense_minutes_pdf_generation(self):
        import datetime
        from defense.minutes.models import DefenseMinutes, MinutesPanelistComment
        from defense.minutes.pdf_generator import generate_minutes_pdf
        from defense.scheduler.models import DefenseSchedule

        schedule = DefenseSchedule.objects.create(
            scope=DefenseSchedule.SCOPE_CAPSTONE,
            semester=self.semester,
            team=self.team_other,
            defense_stage=self.stage,
            scheduled_date=datetime.date(2026, 7, 4),
            start_time=datetime.time(10, 0),
            room='Room 302',
            created_by=self.admin,
        )

        minutes = DefenseMinutes.objects.create(
            schedule=schedule,
            team_name='Test Team Minutes',
            project_title='Automated Test Defense Project',
            adviser_name='Dr. Test Adviser',
            defense_stage_label='Concept Proposal',
            defense_date=datetime.date(2026, 7, 4),
            defense_time=datetime.time(10, 0),
            documenter_name='Test Documenter',
            room='Room 302',
        )
        MinutesPanelistComment.objects.create(
            minutes=minutes,
            panelist_name_snapshot='Dr. Panel Chair',
            panelist_role_snapshot='Chairman',
            comments='Great presentation, proceed to manuscript finalization.',
            display_order=1,
        )

        pdf_bytes = generate_minutes_pdf(minutes)
        self.assertIsInstance(pdf_bytes, bytes)
        self.assertTrue(len(pdf_bytes) > 1000)
        self.assertTrue(pdf_bytes.startswith(b'%PDF'))

    def test_individual_grade_report_generation(self):
        self.client.force_authenticate(user=self.admin)
        from authentication_access_control.models import SystemAuditLog

        url = f'/api/reports/individual-grade/{self.student.id}/'
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertIn('attachment', response['Content-Disposition'])

        # Verify audit trail was logged
        log = SystemAuditLog.objects.filter(action='report.generate_individual_grade').first()
        self.assertIsNotNone(log)
        self.assertEqual(log.category, 'compliance')
    def test_weekly_progress_pdf_generation(self):
        from repository.deliverables.pdf_generator import generate_weekly_reports_pdf

        # Test with empty reports
        pdf_bytes = generate_weekly_reports_pdf(self.team_advised, [])
        self.assertIsInstance(pdf_bytes, bytes)
        self.assertTrue(len(pdf_bytes) > 1000)
        self.assertTrue(pdf_bytes.startswith(b'%PDF'))


